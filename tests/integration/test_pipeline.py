import os
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
import asyncio
import sys
import uuid
import json

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy import text, create_engine
from sqlalchemy.pool import NullPool

# Get test db url (async for app, sync for setup)
TEST_DB_URL = "postgresql+psycopg://postgres:postgres@localhost:5432/gst_copilot_test"
SYNC_TEST_DB_URL = "postgresql+psycopg://postgres:postgres@localhost:5432/gst_copilot_test"

os.environ["DATABASE_URL"] = TEST_DB_URL

from backend.gst_copilot.api.main import app
from backend.gst_copilot.db.database import get_db, Base
from backend.gst_copilot.api.dependencies import get_current_user
from backend.gst_copilot.db.models import User, Client, Organization
from backend.gst_copilot.config import settings

sync_engine = create_engine(SYNC_TEST_DB_URL)

@pytest_asyncio.fixture
async def engine():
    test_engine = create_async_engine(TEST_DB_URL, echo=False, poolclass=NullPool)
    yield test_engine
    await test_engine.dispose()

@pytest_asyncio.fixture
async def test_db_session(engine):
    TestingSessionLocal = async_sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    return TestingSessionLocal

@pytest_asyncio.fixture
async def mock_user_session():
    current_org_id = uuid.uuid4()
    current_user_id = uuid.uuid4()
    
    # Insert the mock org and user into the DB
    with sync_engine.begin() as conn:
        conn.execute(text("INSERT INTO organizations (id, name, created_at) VALUES (:id, :name, NOW())"), {"id": current_org_id, "name": "Demo Org"})
        conn.execute(text("INSERT INTO users (id, email, organization_id, created_at) VALUES (:id, :email, :org_id, NOW())"), {"id": current_user_id, "email": "test@demo.com", "org_id": current_org_id})
    
    async def mock_get_current_user():
        return User(id=current_user_id, email="test@demo.com", organization_id=current_org_id, role="owner")
        
    return mock_get_current_user

@pytest_asyncio.fixture
async def client(test_db_session, mock_user_session):
    async def override_get_db():
        async with test_db_session() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = mock_user_session
    
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
    
    app.dependency_overrides.clear()


@pytest.fixture(autouse=True)
def setup_db():
    assert sync_engine.url.database == "gst_copilot_test", "Refusing to reset a non-test database"
    Base.metadata.drop_all(bind=sync_engine)
    Base.metadata.create_all(bind=sync_engine)
    yield

@pytest.mark.asyncio
async def test_raw_query(test_db_session, mock_user_session):
    async with test_db_session() as session:
        org_id = uuid.uuid4()
        org = Organization(id=org_id, name="Raw Org")
        session.add(org)
        await session.commit()
        
        client = Client(name="Raw Test Client", organization_id=org_id)
        session.add(client)
        await session.commit()
        
        assert client.id is not None

@pytest.mark.asyncio
async def test_full_reconciliation_pipeline(client):
    # 1. Create client
    response = await client.post("/clients", json={"name": "Test Client"})
    assert response.status_code == 200
    client_id = response.json()["id"]

    # Create registration
    response = await client.post(f"/clients/{client_id}/registrations", json={"gstin": "29AAACF1234Z1Z5"})
    assert response.status_code == 200
    reg_id = response.json()["id"]

    # Create period
    response = await client.post(f"/registrations/{reg_id}/periods", json={"period_code": "2026-08"})
    assert response.status_code == 200
    period_id = response.json()["id"]

    # 2. Upload purchase and statement files
    purchase_path = "sample_data/v1/purchase_register.csv"
    statement_path = "sample_data/v1/gstr2b_demo.csv"

    with open(purchase_path, "rb") as f:
        p_resp = await client.post(
            f"/periods/{period_id}/imports", 
            params={"source_type": "purchase"}, 
            files={"file": f}
        )
    assert p_resp.status_code == 200
    p_batch_id = p_resp.json()["id"]
    assert p_resp.json()["status"] == "preview"
    assert p_resp.json()["invalid_count"] == 3

    with open(statement_path, "rb") as f:
        s_resp = await client.post(
            f"/periods/{period_id}/imports", 
            params={"source_type": "statement"}, 
            files={"file": f}
        )
    assert s_resp.status_code == 200
    s_batch_id = s_resp.json()["id"]
    assert s_resp.json()["status"] == "preview"
    assert s_resp.json()["invalid_count"] == 1

    # 4. Acknowledge invalid rows -> commit
    p_com_resp = await client.post(f"/imports/{p_batch_id}/commit", params={"acknowledge_invalid": True, "note": "Synthetic invalid rows reviewed"})
    assert p_com_resp.status_code == 200
    assert p_com_resp.json()["status"] == "committed"
    
    s_com_resp = await client.post(f"/imports/{s_batch_id}/commit", params={"acknowledge_invalid": True, "note": "Synthetic invalid rows reviewed"})
    assert s_com_resp.status_code == 200
    assert s_com_resp.json()["status"] == "committed"

    # 5. Reconcile
    run_resp = await client.post(f"/periods/{period_id}/reconciliation-runs", json={
        "purchase_batch_id": p_batch_id,
        "statement_batch_id": s_batch_id
    })
    assert run_resp.status_code == 200
    run_id = run_resp.json()["id"]
    run_data = run_resp.json()
        
    assert run_data["status"] == "succeeded"
    assert run_data["summary_data"]["distinct_results_by_status"]["matched"] == 7

    summary_data = run_data["summary_data"]
    with open("tests/fixtures/reconciliation_v1/expected_summary.json", "r") as f:
        expected = json.load(f)
        
    assert summary_data["purchase_rows_total"] == expected["purchase_rows_total"]
    assert summary_data["statement_rows_total"] == expected["statement_rows_total"]
    assert summary_data["purchase_rows_valid"] == expected["purchase_rows_valid"]
    assert summary_data["purchase_rows_invalid"] == expected["purchase_rows_invalid"]
    assert summary_data["statement_rows_valid"] == expected["statement_rows_valid"]
    assert summary_data["statement_rows_invalid"] == expected["statement_rows_invalid"]
    assert summary_data["matched_pair_count"] == expected["matched_pair_count"]
    
    for k, v in expected["distinct_results_by_status"].items():
        assert summary_data["distinct_results_by_status"][k] == v

    # Result serialization and evidence must agree with the engine.
    results_url = f"/periods/{period_id}/reconciliation-runs/{run_id}/results"
    response = await client.get(results_url)
    assert response.status_code == 200
    results = response.json()
    assert sum(r["status"] == "matched" for r in results) == 7
    assert all(isinstance(r["purchase_record_ids"], list) for r in results)
    mismatch = next(r for r in results if r["status"] == "amount_mismatch")
    assert mismatch["differences"] and mismatch["purchase_records"] and mismatch["statement_records"]
    resolve_url = f"/runs/{run_id}/results/{mismatch['result_id']}/resolve"
    assert (await client.post(resolve_url, json={"decision": "explained", "note": "  "})).status_code == 400
    for decision in ["explained", "unresolved"]:
        assert (await client.post(resolve_url, json={"decision": decision, "note": "=Synthetic review"})).status_code == 200
    refreshed = (await client.get(results_url)).json()
    updated = next(r for r in refreshed if r["result_id"] == mismatch["result_id"])
    assert updated["review_status"] == "unresolved"
    assert len(updated["history"]) == 2
    assert updated["status"] == "amount_mismatch"
    # Reject imports belonging to another period.
    other = (await client.post(f"/registrations/{reg_id}/periods", json={"period_code": "2026-09"})).json()["id"]
    assert (await client.post(f"/periods/{other}/reconciliation-runs", json={"purchase_batch_id": p_batch_id, "statement_batch_id": s_batch_id})).status_code == 404

    # EXPORT TEST
    export_resp = await client.get(f"/periods/{period_id}/reconciliation-runs/{run_id}/export")
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(export_resp.content))
    assert len(workbook.sheetnames) == 8
    assert workbook["Purchase Records"].max_row == 21
    assert workbook["Statement Records"].max_row == 21
    assert workbook["Review History"].max_row == 3
    assert not any(c.data_type == 'f' for ws in workbook for row in ws for c in row)
    assert any("PUR-" in str(c.value) for row in workbook["Reconciliation Results"] for c in row)

    # Cross Org Export Test
    app.dependency_overrides[get_current_user] = lambda: User(
        id=uuid.uuid4(),
        email="user@org2.com",
        organization_id=uuid.uuid4()
    )
    cross_resp = await client.get(f"/periods/{period_id}/reconciliation-runs/{run_id}/export")
    assert cross_resp.status_code == 404
    app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_cross_org_isolation(client):
    # org1 is already set in the fixture
    resp = await client.post("/clients", json={"name": "Org 1 Client"})
    client_id = resp.json()["id"]
    
    # Run as org2
    org2 = uuid.uuid4()
    with sync_engine.begin() as conn:
        conn.execute(text("INSERT INTO organizations (id, name, created_at) VALUES (:id, :name, NOW())"), {"id": org2, "name": "Org 2"})
        conn.execute(text("INSERT INTO users (id, email, organization_id, created_at) VALUES (:id, :email, :org_id, NOW())"), {"id": uuid.uuid4(), "email": "org2@demo.com", "org_id": org2})

    async def mock_org2_user():
        return User(id=uuid.uuid4(), email="org2@demo.com", organization_id=org2)
    app.dependency_overrides[get_current_user] = mock_org2_user
    resp = await client.get("/clients")
    assert len(resp.json()) == 0
    
    # Try to add reg to org1's client as org2
    resp = await client.post(f"/clients/{client_id}/registrations", json={"gstin": "29AAACF1234Z1Z5"})
    assert resp.status_code == 404

@pytest.mark.asyncio
async def test_idempotent_commit(client):
    client_resp = await client.post("/clients", json={"name": "Test Client"})
    client_id = client_resp.json()["id"]
    reg_resp = await client.post(f"/clients/{client_id}/registrations", json={"gstin": "29AAACF1234Z1Z5"})
    reg_id = reg_resp.json()["id"]
    period_resp = await client.post(f"/registrations/{reg_id}/periods", json={"period_code": "2026-08"})
    period_id = period_resp.json()["id"]
    
    with open("sample_data/v1/purchase_register.csv", "rb") as f:
        p_resp = await client.post(
            f"/periods/{period_id}/imports", 
            params={"source_type": "purchase"}, 
            files={"file": f}
        )
    p_batch_id = p_resp.json()["id"]
    
    # Commit once
    r1 = await client.post(f"/imports/{p_batch_id}/commit", params={"acknowledge_invalid": True, "note": "Synthetic invalid rows reviewed"})
    assert r1.status_code == 200
    
    # Commit again (idempotent)
    r2 = await client.post(f"/imports/{p_batch_id}/commit", params={"acknowledge_invalid": True, "note": "Synthetic invalid rows reviewed"})
    assert r2.status_code == 200
    
@pytest.mark.asyncio
async def test_unacknowledged_invalid_blocks_commit(client):
    client_resp = await client.post("/clients", json={"name": "Test Client"})
    client_id = client_resp.json()["id"]
    reg_resp = await client.post(f"/clients/{client_id}/registrations", json={"gstin": "29AAACF1234Z1Z5"})
    reg_id = reg_resp.json()["id"]
    period_resp = await client.post(f"/registrations/{reg_id}/periods", json={"period_code": "2026-08"})
    period_id = period_resp.json()["id"]
    
    with open("sample_data/v1/purchase_register.csv", "rb") as f:
        p_resp = await client.post(
            f"/periods/{period_id}/imports", 
            params={"source_type": "purchase"}, 
            files={"file": f}
        )
    p_batch_id = p_resp.json()["id"]
    
    # Attempt commit without acknowledging invalid rows
    r1 = await client.post(f"/imports/{p_batch_id}/commit", params={"acknowledge_invalid": False})
    assert r1.status_code == 400
    assert "Acknowledge them to commit" in r1.json()["detail"]


@pytest.mark.asyncio
async def test_category_review_and_mapping_lifecycle(client, mock_user_session):
    from pathlib import Path
    from sqlalchemy import select
    from backend.gst_copilot.db.models import ImportRecord
    c = (await client.post('/clients', json={'name':'Synthetic category client'})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin':'29AAACF1234Z1Z5'})).json()['id']
    period = (await client.post(f'/registrations/{reg}/periods', json={'period_code':'2026-08'})).json()['id']
    # Separate in-memory synthetic description; original fixtures and oracle stay intact.
    payload = Path('sample_data/v1/purchase_register.csv').read_bytes().replace(b'Scenario 1', b'corrugated cartons', 1)
    batch = (await client.post(f'/periods/{period}/imports', params={'source_type':'purchase'}, files={'file':('synthetic.csv',payload,'text/csv')})).json()['id']
    with sync_engine.connect() as conn:
        records = conn.execute(select(ImportRecord.id,ImportRecord.is_valid,ImportRecord.raw_data).where(ImportRecord.batch_id == uuid.UUID(batch)).order_by(ImportRecord.row_number)).all()
    rid = str(records[0].id)
    original = records[0].raw_data
    assert (await client.post(f'/records/{rid}/category-suggestion')).status_code == 400
    assert (await client.post(f'/imports/{batch}/commit',params={'acknowledge_invalid':True,'note':'Synthetic review'})).status_code == 200
    invalid = str(next(r.id for r in records if not r.is_valid))
    assert (await client.post(f'/records/{invalid}/category-suggestion')).status_code == 400
    response = await client.post(f'/records/{rid}/category-suggestion')
    assert response.status_code == 200, response.text
    proposal = response.json()
    assert proposal['source'] == 'mock'
    assert proposal['category'] == 'Packaging'
    assert (await client.post(f'/records/{rid}/category-suggestion')).json()['id'] == proposal['id']
    endpoint = f'/records/{rid}/category-decisions'
    manual = {'action':'manual','category':'Packaging','note':'Synthetic cartons reviewed','save_rule':True,'effective_from':'2026-08-01','previous_decision_id':None}
    response = await client.post(endpoint,json=manual)
    assert response.status_code == 200, response.text
    decision = response.json()
    assert (await client.post(endpoint,json=manual)).status_code == 409
    listing = (await client.get(f'/periods/{period}/categories',params={'batch_id':batch})).json()
    assert listing['items'][0]['category'] == 'Packaging'
    rule = listing['rules'][0]
    assert rule['active']
    mapped = (await client.post(f'/records/{rid}/category-suggestion')).json()
    assert mapped['source'] == 'approved_rule'
    assert (await client.post(f"/category-rules/{rule['id']}/deactivate")).status_code == 200
    accept = {'action':'accept','proposal_id':mapped['id'],'previous_decision_id':decision['id'],'note':'Review again'}
    assert (await client.post(endpoint,json=accept)).status_code == 409
    refreshed = (await client.post(f'/records/{rid}/category-suggestion')).json()
    assert refreshed['source'] == 'mock'
    assert refreshed['id'] != mapped['id']
    reject = dict(accept, action='reject', proposal_id=refreshed['id'])
    response = await client.post(endpoint,json=reject)
    assert response.status_code == 200, response.text
    assert response.json()['category'] == 'Packaging'
    assert (await client.post(endpoint,json=reject)).status_code == 409
    listing = (await client.get(f'/periods/{period}/categories',params={'batch_id':batch})).json()
    assert len(listing['items'][0]['history']) == 2
    with sync_engine.connect() as conn:
        assert conn.execute(select(ImportRecord.raw_data).where(ImportRecord.id == uuid.UUID(rid))).scalar_one() == original
    # A future-dated approved mapping must not apply to this older invoice.
    last = listing['items'][0]['last_decision_id']
    future = dict(manual, previous_decision_id=last, effective_from='2026-09-01')
    future_response = await client.post(endpoint,json=future)
    assert future_response.status_code == 200, future_response.text
    before_effective = (await client.post(f'/records/{rid}/category-suggestion')).json()
    assert before_effective['source'] == 'mock'
    # Accept takes the proposal category, ignoring a caller-supplied replacement.
    accepted = await client.post(endpoint,json={'action':'accept','proposal_id':before_effective['id'],'previous_decision_id':future_response.json()['id'],'category':'Transport','note':'Synthetic suggestion confirmed'})
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()['category'] == 'Packaging'
    # Same organization, separate client: mapping must not leak across clients.
    other = (await client.post('/clients',json={'name':'Other synthetic client'})).json()['id']
    other_reg = (await client.post(f'/clients/{other}/registrations',json={'gstin':'DEMO-OTHER'})).json()['id']
    other_period = (await client.post(f'/registrations/{other_reg}/periods',json={'period_code':'2026-08'})).json()['id']
    other_batch = (await client.post(f'/periods/{other_period}/imports',params={'source_type':'purchase'},files={'file':('synthetic.csv',payload,'text/csv')})).json()['id']
    await client.post(f'/imports/{other_batch}/commit',params={'acknowledge_invalid':True,'note':'Synthetic records reviewed'})
    other_listing = (await client.get(f'/periods/{other_period}/categories',params={'batch_id':other_batch})).json()
    assert other_listing['rules'] == []
    assert other_listing['items'][0]['category'] is None
    # Statement rows are excluded, even after commit.
    statement = (await client.post(f'/periods/{period}/imports',params={'source_type':'statement'},files={'file':('statement.csv',Path('sample_data/v1/gstr2b_demo.csv').read_bytes(),'text/csv')})).json()['id']
    await client.post(f'/imports/{statement}/commit',params={'acknowledge_invalid':True,'note':'Synthetic statement reviewed'})
    with sync_engine.connect() as conn:
        statement_rid = conn.execute(select(ImportRecord.id).where(ImportRecord.batch_id == uuid.UUID(statement))).scalars().first()
    assert (await client.post(f'/records/{statement_rid}/category-suggestion')).status_code == 400
    user = await mock_user_session()
    async def foreign_user():
        return User(id=user.id,email=user.email,organization_id=uuid.uuid4())
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.post(f'/records/{rid}/category-suggestion')).status_code == 404
        assert (await client.post(endpoint,json=manual)).status_code == 404
        assert (await client.get(f'/periods/{period}/categories',params={'batch_id':batch})).status_code == 404
        assert (await client.post(f"/category-rules/{rule['id']}/deactivate")).status_code == 404
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session


@pytest.mark.asyncio
async def test_sales_import_review_export_and_version_isolation(client,mock_user_session):
    from pathlib import Path
    from io import BytesIO
    from openpyxl import load_workbook
    from sqlalchemy import select
    from backend.gst_copilot.db.models import SalesBatch
    name=(await client.post('/clients',json={'name':'Synthetic sales client'})).json()['id']
    reg=(await client.post(f'/clients/{name}/registrations',json={'gstin':'DEMO-SALES'})).json()['id']
    period=(await client.post(f'/registrations/{reg}/periods',json={'period_code':'2026-08'})).json()['id']
    payload=Path('sample_data/sales_v1/sales_register.csv').read_bytes()
    upload_url=f'/periods/{period}/sales-imports'
    assert (await client.post(upload_url,files={'file':('bad.csv',b'wrong,header\n1,2')})).status_code==400
    response=await client.post(upload_url,files={'file':('sales.csv',payload)})
    assert response.status_code==200,response.text
    batch=response.json()['id'];url=f'/sales-imports/{batch}'
    assert (await client.post(upload_url,files={'file':('same.csv',payload)})).json()['id']==batch
    assert len((await client.get(upload_url)).json())==1
    detail=(await client.get(url)).json()
    assert detail['summary']['source_rows']==19 and detail['summary']['included']==0
    first=detail['items'][0];review_url=f"{url}/rows/{first['id']}/review"
    assert (await client.post(review_url,json={'decision':'reviewed','note':'Review'})).status_code==400
    assert (await client.get(url+'/export')).status_code==400
    assert (await client.post(url+'/commit',json={})).status_code==400
    assert (await client.post(url+'/commit',json={'acknowledge_blocked':True,'note':'   '})).status_code==400
    assert (await client.post(url+'/commit',json={'acknowledge_blocked':True,'note':'Synthetic blocked rows acknowledged'})).status_code==200
    assert (await client.post(url+'/commit',json={})).status_code==200
    assert (await client.post(review_url,json={'decision':'reviewed','note':' '})).status_code==400
    assert (await client.post(review_url,json={'decision':'bogus','note':'x'})).status_code==422
    for row in detail['items']:
        endpoint=f"{url}/rows/{row['id']}/review"
        if row['validation_status']!='ready':
            assert (await client.post(endpoint,json={'decision':'reviewed','note':'Cannot override validation'})).status_code==400
        result=await client.post(endpoint,json={'decision':'reviewed' if row['validation_status']=='ready' else 'excluded','note':'Synthetic source reviewed\x01literal'})
        assert result.status_code==200,result.text
    result=(await client.get(url)).json();summary=result['summary']
    expected=json.loads(Path('tests/fixtures/sales_v1/expected_summary.json').read_text())
    assert summary['included_totals']==expected['reviewed_totals']
    assert summary['included']==5 and summary['excluded']==14 and summary['pending']==0
    assert len((await client.get(url,params={'offset':1,'limit':2})).json()['items'])==2
    assert (await client.get(url,params={'status':'duplicate'})).json()['total']==4
    assert (await client.get(url,params={'offset':-1})).status_code==422
    assert (await client.post(review_url,json={'decision':'unresolved','note':'Stale review'})).status_code==409
    newest=result['items'][0]['previous_review_id']
    assert (await client.post(review_url,json={'decision':'unresolved','note':'Reopen for checking','previous_review_id':newest})).status_code==200
    reopened=(await client.get(url)).json()
    assert reopened['summary']['included']==4 and reopened['summary']['pending']==1
    assert reopened['summary']['included_totals']['invoice_total']=='885.00'
    assert len(reopened['items'][0]['history'])==2
    response=await client.get(url+'/export')
    assert response.status_code==200,response.text
    workbook=load_workbook(BytesIO(response.content))
    assert workbook.sheetnames==['Metadata','Summary','Included Totals','Sales Rows','Review History']
    assert workbook['Sales Rows'].max_row==20 and workbook['Review History'].max_row==21
    assert dict(workbook['Included Totals'].values)['invoice_total']=='885.00'
    assert dict(workbook['Summary'].values)['pending']=='1'
    assert not any(c.data_type=='f' for ws in workbook for row in ws for c in row)
    assert any('\\u0001' in str(c.value) for row in workbook['Review History'] for c in row)
    with sync_engine.connect() as conn:
        stored=conn.execute(select(SalesBatch.original_csv).where(SalesBatch.id==uuid.UUID(batch))).scalar_one()
        assert stored.encode('utf-8')==payload
    # A corrected upload has independent review state, and never joins old totals.
    newer=(await client.post(upload_url,files={'file':('corrected.csv',payload.replace(b'S-001',b'NEW-001',1))})).json()['id']
    assert newer!=batch
    assert (await client.get(f'/sales-imports/{newer}')).json()['summary']['included']==0
    assert (await client.get(url)).json()['summary']==reopened['summary']
    # Concurrent identical uploads serialize on the period and return one version.
    concurrent_payload=payload.replace(b'S-001',b'CONCURRENT-001',1)
    attempts=await asyncio.gather(*[client.post(upload_url,files={'file':('concurrent.csv',concurrent_payload)}) for _ in range(2)])
    assert all(r.status_code==200 for r in attempts)
    assert attempts[0].json()['id']==attempts[1].json()['id']
    # Concurrent reviewers with the same previous ID cannot both append a decision.
    current_row=(await client.get(url)).json()['items'][0]
    attempts=await asyncio.gather(*[client.post(review_url,json={'decision':'unresolved','note':'Concurrent review','previous_review_id':current_row['previous_review_id']}) for _ in range(2)])
    assert sorted(r.status_code for r in attempts)==[200,409]
    assert (await client.post(f"/sales-imports/{newer}/rows/{first['id']}/review",json={'decision':'reviewed','note':'Wrong version'})).status_code==400
    assert (await client.post(f'/sales-imports/{newer}/commit',json={'acknowledge_blocked':True,'note':'Review separate version'})).status_code==200
    assert (await client.post(f"/sales-imports/{newer}/rows/{first['id']}/review",json={'decision':'reviewed','note':'Wrong version'})).status_code==404
    # All sales endpoints are scoped, including export and review.
    user=await mock_user_session()
    async def foreign_user(): return User(id=user.id,email=user.email,organization_id=uuid.uuid4())
    app.dependency_overrides[get_current_user]=foreign_user
    try:
        assert (await client.get(upload_url)).status_code==404
        assert (await client.post(upload_url,files={'file':('sales.csv',payload)})).status_code==404
        assert (await client.get(url)).status_code==404
        assert (await client.get(url+'/export')).status_code==404
        assert (await client.post(url+'/commit',json={})).status_code==404
        assert (await client.post(review_url,json={'decision':'reviewed','note':'Unauthorized'})).status_code==404
    finally: app.dependency_overrides[get_current_user]=mock_user_session


async def _period(client, name):
    c = (await client.post('/clients', json={'name': name})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin': 'DEMO-' + name[:8].upper().replace(' ', '')})).json()['id']
    return (await client.post(f'/registrations/{reg}/periods', json={'period_code': '2026-08'})).json()['id']

@pytest.mark.asyncio
async def test_validation_evidence_uses_physical_row_and_rejected_uploads_leave_no_file(client):
    from pathlib import Path
    period = await _period(client, 'Blank id client')
    lines = Path('sample_data/v1/purchase_register.csv').read_text().splitlines()
    # Two rows with blank record_id must each show only their own row as evidence.
    lines[1] = ',' + lines[1].split(',', 1)[1]
    lines[2] = ',' + lines[2].split(',', 1)[1]
    payload = ('\n'.join(lines) + '\n').encode()
    p = (await client.post(f'/periods/{period}/imports', params={'source_type': 'purchase'}, files={'file': ('p.csv', payload)})).json()['id']
    s = (await client.post(f'/periods/{period}/imports', params={'source_type': 'statement'}, files={'file': ('s.csv', Path('sample_data/v1/gstr2b_demo.csv').read_bytes())})).json()['id']
    for b in (p, s):
        assert (await client.post(f'/imports/{b}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})).status_code == 200
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': p, 'statement_batch_id': s})).json()['id']
    results = (await client.get(f'/periods/{period}/reconciliation-runs/{run}/results')).json()
    blank = [r for r in results if r['status'] == 'validation_error' and r['purchase_record_ids'] == ['']]
    assert len(blank) == 2
    assert sorted(r['purchase_records'][0]['row_number'] for r in blank) == [2, 3]
    assert all(len(r['purchase_records']) == 1 for r in blank)
    # Duplicate record IDs reject the file; its stored copy must be removed.
    storage = Path(settings.STORAGE_DIR)
    before = sum(1 for _ in storage.rglob('*.csv')) if storage.exists() else 0
    dup = ('\n'.join(lines[:2] + [lines[3], lines[3]]) + '\n').encode()
    assert (await client.post(f'/periods/{period}/imports', params={'source_type': 'purchase'}, files={'file': ('dup.csv', dup)})).status_code == 400
    assert sum(1 for _ in storage.rglob('*.csv')) == before


@pytest.mark.asyncio
async def test_tax_worksheet_approval_lifecycle(client, mock_user_session):
    from pathlib import Path
    from io import BytesIO
    from openpyxl import load_workbook
    reference = json.loads(Path('tests/fixtures/calc_v1/cases.json').read_text())['integration_case']['expected']
    period = await _period(client, 'Worksheet client')
    # Sales: review every ready row, exclude the rest.
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('sales.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()['id']
    assert (await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})).status_code == 200
    rows = (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']
    for row in rows:
        decision = 'reviewed' if row['validation_status'] == 'ready' else 'excluded'
        assert (await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': decision, 'note': 'Synthetic'})).status_code == 200
    # Purchases: reconcile the v1 pack.
    batches = {}
    for kind, path in (('purchase', 'sample_data/v1/purchase_register.csv'), ('statement', 'sample_data/v1/gstr2b_demo.csv')):
        batches[kind] = (await client.post(f'/periods/{period}/imports', params={'source_type': kind}, files={'file': (kind, Path(path).read_bytes())})).json()['id']
        await client.post(f'/imports/{batches[kind]}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': batches['purchase'], 'statement_batch_id': batches['statement']})).json()['id']
    preview_url = f'/periods/{period}/worksheet'
    params = {'sales_batch_id': sales, 'run_id': run}
    first = (await client.get(preview_url, params=params)).json()
    assert any('no ITC decision' in b for b in first['blockers'])

    itc_url = f'/runs/{run}/itc-decisions'
    results = (await client.get(itc_url)).json()
    assert len(results) == 25 and all(r['decision'] == 'undecided' for r in results)
    target = next(r for r in results if r['purchase_record_ids'] == ['PUR-001'])
    mismatch = next(r for r in results if r['status'] == 'amount_mismatch')
    assert (await client.post(itc_url, json={'items': [{'result_id': mismatch['result_id'], 'decision': 'claim'}], 'note': 'x'})).status_code == 400
    others = [{'result_id': r['result_id'], 'decision': 'not_claimed'} for r in results if r is not target]
    assert (await client.post(itc_url, json={'items': others + [{'result_id': target['result_id'], 'decision': 'claim'}], 'note': 'Synthetic ITC review'})).status_code == 200
    # Saving again against the old (empty) previous decision is stale.
    assert (await client.post(itc_url, json={'items': [{'result_id': target['result_id'], 'decision': 'claim'}], 'note': 'again'})).status_code == 409

    adj_url = f'/periods/{period}/adjustments'
    assert (await client.post(adj_url, json={'adjustment_type': 'rcm_liability', 'tax_head': 'cgst', 'amount': '5', 'note': 'x'})).status_code == 400
    assert (await client.post(adj_url, json={'adjustment_type': 'interest', 'tax_head': 'cgst', 'amount': '5.00', 'note': 'x'})).status_code == 422
    assert (await client.post(adj_url, json={'adjustment_type': 'rcm_liability', 'tax_head': 'cgst', 'amount': '5.00', 'note': 'Synthetic RCM'})).status_code == 200
    wrong = (await client.post(adj_url, json={'adjustment_type': 'other_liability', 'tax_head': 'igst', 'amount': '1000.00', 'note': 'Entered by mistake'})).json()['id']
    assert (await client.post(f'/adjustments/{wrong}/void', json={'reason': 'Keyed in error'})).status_code == 200
    assert (await client.post(f'/adjustments/{wrong}/void', json={'reason': 'again'})).status_code == 409

    ready = (await client.get(preview_url, params=params)).json()
    assert ready['blockers'] == []
    for head, values in reference.items():
        for key, value in values.items():
            assert ready['payload']['worksheet']['heads'][head][key] == value, (head, key)

    draft = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    assert draft['state'] == 'draft' and draft['fingerprint'] == ready['fingerprint']
    approve_url = f"/tax-drafts/{draft['id']}/approve"
    assert (await client.post(approve_url, json={'note': ' '})).status_code == 400
    approval = await client.post(approve_url, json={'note': 'Synthetic approval'})
    assert approval.status_code == 200, approval.text
    approval_id = approval.json()['id']
    assert (await client.post(approve_url, json={'note': 'Twice'})).status_code == 409
    drafts_url = f'/periods/{period}/tax-drafts'
    assert (await client.get(drafts_url)).json()[0]['state'] == 'approved'

    # Any input change makes the approval stale; the frozen snapshot keeps its figures.
    ready_row = next(r for r in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items'] if r['decision'] == 'reviewed')
    assert (await client.post(f"/sales-imports/{sales}/rows/{ready_row['id']}/review", json={'decision': 'unresolved', 'note': 'Recheck', 'previous_review_id': ready_row['previous_review_id']})).status_code == 200
    stale = (await client.get(drafts_url)).json()[0]
    assert stale['state'] == 'approved_stale'
    assert stale['payload']['worksheet']['heads']['cgst']['net'] == reference['cgst']['net']
    row = next(b for b in (await client.get('/board')).json() if b['period_id'] == period)
    assert row['worksheet'] == {'tone': 'blocked', 'label': 'Approval out of date', 'target': 'worksheet'}
    assert row['sales']['label'] == '1 sales row(s) to review' and row['purchases']['tone'] == 'done'
    export = await client.get(f"/tax-drafts/{draft['id']}/export")
    assert export.status_code == 200
    wb = load_workbook(BytesIO(export.content))
    assert wb.sheetnames == ['Metadata', 'Worksheet', 'Set-off', 'GSTR-3B view', 'Output Rows', 'ITC Claimed', 'Adjustments', 'Filing Evidence']
    lines = {r[0]: r for r in wb['Worksheet'].iter_rows(min_row=2, values_only=True)}
    assert lines['cgst'][-1] == reference['cgst']['net'] and lines['igst'][3] == reference['igst']['liability']
    assert dict(wb['Metadata'].iter_rows(min_row=2, values_only=True))['state'] == 'approved_stale'
    assert wb['Output Rows'].max_row == 6 and wb['ITC Claimed'].max_row == 2 and wb['Adjustments'].max_row == 2
    assert not any(c.data_type == 'f' for ws in wb for row in ws for c in row)

    evidence_url = f'/tax-approvals/{approval_id}/filing-evidence'
    assert (await client.post(evidence_url, json={'arn': 'DEMO-ARN-1', 'filed_on': '2026-09-20', 'note': 'User reported'})).json()['label'] == 'user-reported, not verified'
    assert (await client.post(f'/tax-approvals/{approval_id}/reopen', json={'reason': ' '})).status_code == 400
    assert (await client.post(f'/tax-approvals/{approval_id}/reopen', json={'reason': 'Sales row reopened'})).status_code == 200
    assert (await client.post(f'/tax-approvals/{approval_id}/reopen', json={'reason': 'again'})).status_code == 409
    assert (await client.post(evidence_url, json={'arn': 'X', 'filed_on': '2026-09-20'})).status_code == 409

    blocked = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    assert blocked['blockers'] and (await client.post(f"/tax-drafts/{blocked['id']}/approve", json={'note': 'x'})).status_code == 400
    ready_row = next(r for r in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items'] if r['id'] == ready_row['id'])
    await client.post(f"/sales-imports/{sales}/rows/{ready_row['id']}/review", json={'decision': 'reviewed', 'note': 'Confirmed', 'previous_review_id': ready_row['previous_review_id']})
    second = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    assert second['blockers'] == []
    # Inputs changed after drafting -> approval refused.
    await client.post(adj_url, json={'adjustment_type': 'other_credit', 'tax_head': 'sgst', 'amount': '1.00', 'note': 'Late credit'})
    assert (await client.post(f"/tax-drafts/{second['id']}/approve", json={'note': 'x'})).status_code == 409
    third = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    assert third['payload']['worksheet']['heads']['sgst']['net'] == '21.50'
    assert (await client.post(f"/tax-drafts/{third['id']}/approve", json={'note': 'Re-approved'})).status_code == 200
    states = [d['state'] for d in (await client.get(drafts_url)).json()]
    assert states == ['approved', 'draft', 'draft', 'reopened']

    actions = {e['action'] for e in (await client.get('/audit-events', params={'period_id': period, 'limit': 200})).json()}
    assert {'sales_uploaded', 'import_uploaded', 'import_committed', 'reconciliation_run', 'itc_decided', 'adjustment_added', 'adjustment_voided',
            'tax_draft_created', 'tax_draft_approved', 'tax_approval_reopened', 'filing_evidence_recorded'} <= actions

    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4())
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.get(preview_url, params=params)).status_code == 404
        assert (await client.get(itc_url)).status_code == 404
        assert (await client.post(adj_url, json={'adjustment_type': 'rcm_liability', 'tax_head': 'cgst', 'amount': '1.00', 'note': 'x'})).status_code == 404
        assert (await client.post(f"/tax-drafts/{third['id']}/approve", json={'note': 'x'})).status_code == 404
        assert (await client.get(f"/tax-drafts/{third['id']}/export")).status_code == 404
        assert (await client.post(f'/tax-approvals/{approval_id}/reopen', json={'reason': 'x'})).status_code == 404
        assert (await client.get('/audit-events', params={'period_id': period})).json() == []
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session


@pytest.mark.asyncio
async def test_concurrent_approvals_and_itc_decisions_serialize(client):
    from pathlib import Path
    period = await _period(client, 'Concurrency client')
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()['id']
    await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    for row in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']:
        await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': 'excluded', 'note': 'Synthetic'})
    b = {}
    for kind, path in (('purchase', 'sample_data/v1/purchase_register.csv'), ('statement', 'sample_data/v1/gstr2b_demo.csv')):
        b[kind] = (await client.post(f'/periods/{period}/imports', params={'source_type': kind}, files={'file': (kind, Path(path).read_bytes())})).json()['id']
        await client.post(f'/imports/{b[kind]}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': b['purchase'], 'statement_batch_id': b['statement']})).json()['id']
    results = (await client.get(f'/runs/{run}/itc-decisions')).json()
    first = results[0]
    # Two reviewers deciding the same result from the same starting point: exactly one wins.
    same = {'items': [{'result_id': first['result_id'], 'decision': 'deferred'}], 'note': 'Concurrent'}
    attempts = await asyncio.gather(*[client.post(f'/runs/{run}/itc-decisions', json=same) for _ in range(2)])
    assert sorted(r.status_code for r in attempts) == [200, 409]
    rest = [{'result_id': r['result_id'], 'decision': 'not_claimed'} for r in results[1:]]
    assert (await client.post(f'/runs/{run}/itc-decisions', json={'items': rest, 'note': 'Synthetic'})).status_code == 200
    draft = (await client.post(f'/periods/{period}/tax-drafts', json={'sales_batch_id': sales, 'run_id': run})).json()
    assert draft['blockers'] == []
    attempts = await asyncio.gather(*[client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'Concurrent'}) for _ in range(2)])
    assert sorted(r.status_code for r in attempts) == [200, 409]
    assert [d['state'] for d in (await client.get(f'/periods/{period}/tax-drafts')).json()] == ['approved']


@pytest.mark.asyncio
async def test_board_reports_status_from_records_and_is_org_scoped(client, mock_user_session):
    from pathlib import Path
    period = await _period(client, 'Board client')
    row = next(b for b in (await client.get('/board')).json() if b['period_id'] == period)
    assert [row[k]['tone'] for k in ('sales', 'purchases', 'worksheet')] == ['waiting', 'waiting', 'waiting']
    await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})
    await client.post(f'/periods/{period}/imports', params={'source_type': 'purchase'}, files={'file': ('p.csv', Path('sample_data/v1/purchase_register.csv').read_bytes())})
    row = next(b for b in (await client.get('/board')).json() if b['period_id'] == period)
    assert row['sales']['label'] == 'Commit sales import' and row['purchases']['label'] == 'Commit imports'
    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4())
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.get('/board')).json() == []
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session


@pytest.mark.asyncio
async def test_assistants_suggest_explain_and_never_write_without_a_person(client, mock_user_session):
    from pathlib import Path
    period = await _period(client, 'Assist client')
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()['id']
    await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    for row in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']:
        await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': 'reviewed' if row['validation_status'] == 'ready' else 'excluded', 'note': 'Synthetic'})
    # Column mapper: a Tally-style header layout converts to the template; values are copied verbatim.
    original = Path('sample_data/v1/purchase_register.csv').read_bytes()
    tally = original.replace(b'record_id,document_type,supplier_ref,invoice_number,invoice_date,taxable_value,cgst,sgst,igst,cess,invoice_total,description',
                             b'Sl No,Vch Type,Party Code,Bill No,Bill Date,Taxable Amt,CGST Amt,SGST Amt,IGST Amt,Cess Amt,Grand Total,Narration', 1)
    assert (await client.post(f'/periods/{period}/imports', params={'source_type': 'purchase'}, files={'file': ('tally.csv', tally)})).status_code == 400
    proposal = (await client.post('/column-mapping/propose', data={'template': 'purchase'}, files={'file': ('tally.csv', tally)})).json()
    assert proposal['missing_required'] == [] and proposal['mapping']['invoice_number'] == 'Bill No'
    converted = await client.post('/column-mapping/apply', data={'template': 'purchase', 'mapping': json.dumps(proposal['mapping'])}, files={'file': ('tally.csv', tally)})
    assert converted.status_code == 200 and converted.content == original.replace(b'\r\n', b'\n')
    bad = await client.post('/column-mapping/apply', data={'template': 'purchase', 'mapping': json.dumps({'record_id': 'Sl No'})}, files={'file': ('tally.csv', tally)})
    assert bad.status_code == 400
    b = {'purchase': (await client.post(f'/periods/{period}/imports', params={'source_type': 'purchase'}, files={'file': ('mapped.csv', converted.content)})).json()}
    assert b['purchase']['invalid_count'] == 3
    b['statement'] = (await client.post(f'/periods/{period}/imports', params={'source_type': 'statement'}, files={'file': ('s.csv', Path('sample_data/v1/gstr2b_demo.csv').read_bytes())})).json()
    for x in b.values(): await client.post(f"/imports/{x['id']}/commit", params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': b['purchase']['id'], 'statement_batch_id': b['statement']['id']})).json()['id']

    # ITC pre-fill: suggestions only; nothing recorded until accepted through the normal endpoint.
    sugg = (await client.get(f'/runs/{run}/itc-suggestions')).json()
    assert sugg['counts'] == {'claim': 7, 'deferred': 3, 'not_claimed': 15}
    assert all(r['decision'] == 'undecided' for r in (await client.get(f'/runs/{run}/itc-decisions')).json())
    items = [{'result_id': s['result_id'], 'decision': s['decision']} for s in sugg['suggestions']]
    assert (await client.post(f'/runs/{run}/itc-decisions', json={'items': items, 'note': 'Accepted suggestions'})).status_code == 200
    assert (await client.get(f'/runs/{run}/itc-suggestions')).json()['suggestions'] == []

    results = (await client.get(f'/periods/{period}/reconciliation-runs/{run}/results')).json()
    books_only = next(r for r in results if r['purchase_record_ids'] == ['PUR-006'])
    inv = (await client.get(f"/runs/{run}/results/{books_only['result_id']}/investigate")).json()
    assert inv['status'] == 'books_only' and inv['records'][0]['record_id'] == 'PUR-006' and inv['suggested_note']
    swap = next(r for r in results if r['purchase_record_ids'] == ['PUR-012'])
    assert any('inter-state' in f for f in (await client.get(f"/runs/{run}/results/{swap['result_id']}/investigate")).json()['findings'])
    assert (await client.get(f'/runs/{run}/results/RES-NOPE/investigate')).status_code == 404

    drafts = (await client.get(f'/runs/{run}/supplier-followups')).json()['drafts']
    assert {d['supplier_ref'] for d in drafts} == {'DEMO-SUP-001', 'DEMO-SUP-002', 'DEMO-SUP-003'}  # books-only 006/007/016, mismatches 010/011/012
    assert any('INV-106' in d['message'] for d in drafts)

    # Knowledge rule: drafted from a note, confirmed by a person, applied once to the period.
    client_id = (await client.get('/board')).json()
    client_id = next(r['client_id'] for r in client_id if r['period_id'] == period)
    rule = (await client.post('/knowledge-rules/draft', json={'client_id': client_id, 'note': 'Rent to unregistered landlord: RCM ₹4,500 CGST monthly from Aug 2026'})).json()
    assert rule['status'] == 'proposed' and rule['amount'] == '4500.00' and rule['missing'] == []
    assert (await client.post(f"/knowledge-rules/{rule['id']}/apply", json={'period_id': period})).status_code == 409
    confirm = {'rule_kind': 'adjustment', 'adjustment_type': 'rcm_liability', 'tax_head': 'cgst', 'amount': '4500.00', 'effective_from': '2026-08'}
    assert (await client.post(f"/knowledge-rules/{rule['id']}/confirm", json=confirm)).json()['status'] == 'active'
    assert (await client.post(f"/knowledge-rules/{rule['id']}/confirm", json=confirm)).status_code == 409
    assert [r['applied'] for r in (await client.get(f'/periods/{period}/applicable-rules')).json()] == [False]
    assert (await client.post(f"/knowledge-rules/{rule['id']}/apply", json={'period_id': period})).status_code == 200
    assert (await client.post(f"/knowledge-rules/{rule['id']}/apply", json={'period_id': period})).status_code == 409
    other_period = await _period(client, 'Other rule client')
    assert (await client.post(f"/knowledge-rules/{rule['id']}/apply", json={'period_id': other_period})).status_code == 400
    note_only = (await client.post('/knowledge-rules/draft', json={'client_id': client_id, 'note': 'Client pays some fees'})).json()
    assert set(note_only['missing']) == {'adjustment_type', 'tax_head', 'amount', 'effective_from'}
    assert (await client.post(f"/knowledge-rules/{note_only['id']}/dismiss")).json()['status'] == 'dismissed'

    # Savings finder: figures come from recorded rows (hand-checked against the fixture).
    savings = (await client.get(f'/periods/{period}/savings')).json()
    by_kind = {i['kind']: i for i in savings['items']}
    assert by_kind['waiting_on_supplier']['amount'] == '414.00'
    assert by_kind['statement_shows_more_tax']['amount'] == '20.00'
    assert by_kind['rcm_without_credit']['amount'] == '4500.00' and by_kind['rcm_without_credit']['needs_ca']
    assert savings['total'] == '4934.00'
    assert (await client.get('/savings')).json()['total'] == '4934.00'

    brief = (await client.get('/brief')).json()['items']
    assert any('₹4,934.00' in b['text'] for b in brief) and all(b['href'].startswith('/') for b in brief)

    await client.post(f'/periods/{period}/tax-drafts', json={'sales_batch_id': sales, 'run_id': run})
    ask = lambda q, **kw: client.post(f'/periods/{period}/ask', json={'question': q, **kw})
    head = (await ask('Why is CGST net what it is?')).json()
    assert head['intent'] == 'explain_head' and head['citations']
    assert 'Nothing is blocking' in (await ask('What is blocking approval?', sales_batch_id=sales, run_id=run)).json()['answer']
    assert (await ask('How much credit is at risk?')).json()['answer'].startswith('₹4934.00')
    assert (await ask('What is the GST rate on gold?')).json()['intent'] is None
    actions = {e['action'] for e in (await client.get('/audit-events', params={'limit': 200})).json()}
    assert {'rule_proposed', 'rule_confirmed', 'rule_applied', 'rule_dismissed'} <= actions

    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4())
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.get(f'/runs/{run}/itc-suggestions')).status_code == 404
        assert (await client.get(f"/runs/{run}/results/{books_only['result_id']}/investigate")).status_code == 404
        assert (await client.get(f'/runs/{run}/supplier-followups')).status_code == 404
        assert (await client.get(f'/periods/{period}/savings')).status_code == 404
        assert (await ask('Who approved this period?')).status_code == 404
        assert (await client.post(f"/knowledge-rules/{rule['id']}/retire")).status_code == 404
        assert (await client.get('/knowledge-rules')).json() == []
        assert (await client.get('/savings')).json()['total'] == '0.00'
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session


@pytest.mark.asyncio
async def test_legal_register_reminders_frequencies_and_cache(client, mock_user_session):
    """Arbitrary test values throughout — not statutory rates or dates."""
    from pathlib import Path
    period = await _period(client, 'Rules client')
    board = (await client.get('/board')).json()
    client_id = next(r['client_id'] for r in board if r['period_id'] == period)
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()['id']
    await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    for row in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']:
        await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': 'reviewed' if row['validation_status'] == 'ready' else 'excluded', 'note': 'Synthetic'})
    b = {}
    for kind, path in (('purchase', 'sample_data/v1/purchase_register.csv'), ('statement', 'sample_data/v1/gstr2b_demo.csv')):
        b[kind] = (await client.post(f'/periods/{period}/imports', params={'source_type': kind}, files={'file': (kind, Path(path).read_bytes())})).json()['id']
        await client.post(f'/imports/{b[kind]}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': b['purchase'], 'statement_batch_id': b['statement']})).json()['id']
    as_of = {'as_of': '2026-09-30'}

    # Nothing runs until a rule is confirmed.
    register = (await client.get('/legal-rules')).json()['rules']
    assert {r['key'] for r in register} == {'itc_claim_deadline', 'blocked_credit_keywords', 'late_payment_interest', 'late_fee', 'return_due_dates', 'credit_utilisation_order', 'payment_rounding'} and all(r['active'] is None for r in register)
    confirm = lambda key, value, **kw: client.post(f'/legal-rules/{key}/confirm', json={'value': value, 'source_reference': 'Test fixture value, not law', 'effective_from': '2026-04', 'checked_against_current_law': True} | kw)
    assert (await confirm('blocked_credit_keywords', {'keywords': ['scenario 5']}, checked_against_current_law=False)).status_code == 400
    assert (await confirm('itc_claim_deadline', {'day': 31, 'month': 4, 'warn_days': 10})).status_code == 400
    assert (await confirm('blocked_credit_keywords', {'keywords': ['scenario 5']})).status_code == 200

    # Blocked-credit words change the ITC pre-fill: PUR-005 ("Scenario 5") is deferred, not claimed.
    sugg = (await client.get(f'/runs/{run}/itc-suggestions')).json()
    assert sugg['counts'] == {'claim': 6, 'deferred': 4, 'not_claimed': 15}
    pur5 = next(s for s in sugg['suggestions'] if 'blocked-credit' in s['reason'])
    items = [{'result_id': s['result_id'], 'decision': 'claim' if s is pur5 else s['decision']} for s in sugg['suggestions']]
    assert (await client.post(f'/runs/{run}/itc-decisions', json={'items': items, 'note': 'Claimed PUR-005 deliberately'})).status_code == 200
    got = {i['kind']: i for i in (await client.get(f'/periods/{period}/savings', params=as_of)).json()['items']}
    assert got['possibly_blocked_credit']['amount'] == '900.00' and 'PUR-005' in got['possibly_blocked_credit']['detail']
    assert 'claim_deadline' not in got and 'interest_estimate' not in got

    assert (await confirm('itc_claim_deadline', {'day': 30, 'month': 11, 'warn_days': 500})).status_code == 200
    got = {i['kind']: i for i in (await client.get(f'/periods/{period}/savings', params=as_of)).json()['items']}
    assert got['claim_deadline']['amount'] == '414.00' and '2027-11-30' in got['claim_deadline']['detail']
    assert (await confirm('itc_claim_deadline', {'day': 30, 'month': 11, 'warn_days': 30})).status_code == 200  # new version, retires the old
    reg = next(r for r in (await client.get('/legal-rules')).json()['rules'] if r['key'] == 'itc_claim_deadline')
    assert len(reg['history']) == 2 and reg['active']['value']['warn_days'] == 30
    assert 'claim_deadline' not in {i['kind'] for i in (await client.get(f'/periods/{period}/savings', params=as_of)).json()['items']}

    # Interest and late fee estimates need a draft with a positive net and no filing reference.
    await client.post(f'/periods/{period}/adjustments', json={'adjustment_type': 'other_liability', 'tax_head': 'igst', 'amount': '1000.00', 'note': 'Synthetic liability'})
    await client.post(f'/periods/{period}/tax-drafts', json={'sales_batch_id': sales, 'run_id': run})
    await confirm('late_payment_interest', {'rate_percent': '18', 'due_day': 20})
    await confirm('late_fee', {'per_day': '50', 'cap': '200', 'due_day': 20})
    got = {i['kind']: i for i in (await client.get(f'/periods/{period}/savings', params=as_of)).json()['items']}
    assert got['interest_estimate']['amount'] == '3.60' and got['late_fee_estimate']['amount'] == '200.00'
    assert (await client.post('/legal-rules/late_fee/retire')).status_code == 200
    assert 'late_fee_estimate' not in {i['kind'] for i in (await client.get(f'/periods/{period}/savings', params=as_of)).json()['items']}

    # Reminders: firm-wide, block approval until acknowledged for the period.
    reminder = (await client.post('/knowledge-rules/draft', json={'client_id': None, 'note': 'Confirm cash payment received before filing, from Aug 2026'})).json()
    assert reminder['rule_kind'] == 'reminder' and reminder['client_name'] == 'All clients (firm-wide)'
    assert (await client.post(f"/knowledge-rules/{reminder['id']}/confirm", json={'rule_kind': 'adjustment', 'adjustment_type': 'other_liability', 'tax_head': 'igst', 'amount': '1.00', 'effective_from': '2026-08'})).status_code == 400
    assert (await client.post(f"/knowledge-rules/{reminder['id']}/confirm", json={'rule_kind': 'reminder', 'effective_from': '2026-08'})).status_code == 200
    draft = (await client.post(f'/periods/{period}/tax-drafts', json={'sales_batch_id': sales, 'run_id': run})).json()
    assert any('reminder' in b for b in draft['blockers'])
    assert (await client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'x'})).status_code == 400
    ack = f"/knowledge-rules/{reminder['id']}/acknowledge"
    assert (await client.post(ack, json={'period_id': period, 'note': ' '})).status_code == 400
    assert (await client.post(ack, json={'period_id': period, 'note': 'Bank statement checked'})).status_code == 200
    assert (await client.post(ack, json={'period_id': period, 'note': 'again'})).status_code == 409
    applicable = (await client.get(f'/periods/{period}/applicable-rules')).json()
    assert next(r for r in applicable if r['id'] == reminder['id'])['acknowledged']['note'] == 'Bank statement checked'
    draft = (await client.post(f'/periods/{period}/tax-drafts', json={'sales_batch_id': sales, 'run_id': run})).json()
    assert draft['blockers'] == [] and (await client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'Checked'})).status_code == 200

    # Frequencies: a quarterly rule from June applies in September, not August.
    quarterly = (await client.post('/knowledge-rules/draft', json={'client_id': client_id, 'note': 'Quarterly RCM on legal fees Rs. 900.00 SGST from 2026-06'})).json()
    assert quarterly['frequency'] == 'quarterly'
    await client.post(f"/knowledge-rules/{quarterly['id']}/confirm", json={'rule_kind': 'adjustment', 'frequency': 'quarterly', 'adjustment_type': 'rcm_liability', 'tax_head': 'sgst', 'amount': '900.00', 'effective_from': '2026-06'})
    assert quarterly['id'] not in {r['id'] for r in (await client.get(f'/periods/{period}/applicable-rules')).json()}
    assert (await client.post(f"/knowledge-rules/{quarterly['id']}/apply", json={'period_id': period})).status_code == 400
    reg_id = next(r for r in (await client.get(f'/clients/{client_id}/registrations')).json())['id']
    september = (await client.post(f'/registrations/{reg_id}/periods', json={'period_code': '2026-09'})).json()['id']
    assert (await client.post(f"/knowledge-rules/{quarterly['id']}/apply", json={'period_id': september})).status_code == 200

    # Cache: a new period shows on the Board immediately (every change is audited, which invalidates the cache).
    before = {r['period_id'] for r in (await client.get('/board')).json()}
    october = (await client.post(f'/registrations/{reg_id}/periods', json={'period_code': '2026-10'})).json()['id']
    after = {r['period_id'] for r in (await client.get('/board')).json()}
    assert september in before and october not in before and october in after
    assert {'period_created', 'legal_rule_confirmed', 'legal_rule_retired', 'reminder_acknowledged'} <= {e['action'] for e in (await client.get('/audit-events', params={'limit': 200})).json()}

    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4())
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert all(r['active'] is None for r in (await client.get('/legal-rules')).json()['rules'])
        assert (await client.post('/legal-rules/blocked_credit_keywords/retire')).status_code == 404
        assert (await client.post(ack, json={'period_id': period, 'note': 'x'})).status_code == 404
        assert (await client.get('/knowledge-rules')).json() == []
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session


@pytest.mark.asyncio
async def test_reminder_confirmed_after_draft_still_blocks_approval(client):
    from pathlib import Path
    period = await _period(client, 'Late reminder client')
    client_id = next(r['client_id'] for r in (await client.get('/board')).json() if r['period_id'] == period)
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()['id']
    await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    for row in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']:
        await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': 'reviewed' if row['validation_status'] == 'ready' else 'excluded', 'note': 'Synthetic'})
    b = {}
    for kind, path in (('purchase', 'sample_data/v1/purchase_register.csv'), ('statement', 'sample_data/v1/gstr2b_demo.csv')):
        b[kind] = (await client.post(f'/periods/{period}/imports', params={'source_type': kind}, files={'file': (kind, Path(path).read_bytes())})).json()['id']
        await client.post(f'/imports/{b[kind]}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': b['purchase'], 'statement_batch_id': b['statement']})).json()['id']
    sugg = (await client.get(f'/runs/{run}/itc-suggestions')).json()['suggestions']
    await client.post(f'/runs/{run}/itc-decisions', json={'items': [{'result_id': s['result_id'], 'decision': s['decision']} for s in sugg], 'note': 'Synthetic'})
    draft = (await client.post(f'/periods/{period}/tax-drafts', json={'sales_batch_id': sales, 'run_id': run})).json()
    assert draft['blockers'] == []
    rule = (await client.post('/knowledge-rules/draft', json={'client_id': client_id, 'note': 'Ensure challan is paid before filing, from Aug 2026'})).json()
    await client.post(f"/knowledge-rules/{rule['id']}/confirm", json={'rule_kind': 'reminder', 'effective_from': '2026-08'})
    blocked = await client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'x'})
    assert blocked.status_code == 400 and 'reminder' in blocked.json()['detail']


@pytest_asyncio.fixture
async def anon_client(test_db_session):
    """A client with NO user override: requests go through real session-cookie authentication."""
    from backend.gst_copilot.api.routers.auth import throttle
    throttle.failures.clear()
    async def override_get_db():
        async with test_db_session() as session:
            yield session
    app.dependency_overrides[get_db] = override_get_db
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        yield ac
    app.dependency_overrides.clear()

def _make_user(email, password):
    from backend.gst_copilot.auth import hash_password
    org_id = uuid.uuid4()
    with sync_engine.begin() as conn:
        conn.execute(text("INSERT INTO organizations (id, name, created_at) VALUES (:id, 'GST Helper Demo CA Firm', NOW())"), {"id": org_id})
        conn.execute(text("INSERT INTO users (id, email, organization_id, created_at, password_hash, is_active) VALUES (:id, :email, :org, NOW(), :ph, true)"),
                     {"id": uuid.uuid4(), "email": email, "org": org_id, "ph": hash_password(password) if password else None})
    return org_id

@pytest.mark.asyncio
async def test_every_route_except_health_and_login_requires_a_session(anon_client):
    # The client upload link is the only public data route: the token in the URL is the credential (tested separately).
    open_paths = {'/health/live', '/health/ready', '/auth/login', '/openapi.json', '/docs', '/docs/oauth2-redirect', '/redoc', '/public/upload/{token}'}
    checked = 0
    for route in app.routes:
        if route.path in open_paths or not hasattr(route, 'methods'): continue
        path = route.path
        for name in route.param_convertors: path = path.replace('{' + name + '}', '00000000-0000-0000-0000-000000000000' if 'id' in name else 'x')
        for method in route.methods - {'HEAD', 'OPTIONS'}:
            response = await anon_client.request(method, path)
            assert response.status_code == 401, (method, route.path, response.status_code)
            checked += 1
    assert checked > 50

@pytest.mark.asyncio
async def test_login_session_logout_lockout_and_expiry(anon_client):
    from backend.gst_copilot.config import settings
    _make_user('admin@demo.com', 'synthetic-test-password')
    _make_user('nopass@demo.com', None)
    assert (await anon_client.get('/clients')).status_code == 401
    assert (await anon_client.post('/auth/login', json={'email': 'admin@demo.com', 'password': 'wrong-password'})).status_code == 401
    assert (await anon_client.post('/auth/login', json={'email': 'nobody@demo.com', 'password': 'whatever-it-is'})).json()['detail'] == 'Email or password is incorrect.'
    assert (await anon_client.post('/auth/login', json={'email': 'nopass@demo.com', 'password': 'anything-here'})).status_code == 401
    ok = await anon_client.post('/auth/login', json={'email': 'ADMIN@demo.com ', 'password': 'synthetic-test-password'})
    body = ok.json()
    assert ok.status_code == 200 and body['email'] == 'admin@demo.com' and body['firm'] == 'GST Helper Demo CA Firm' and body['role'] == 'owner' and body['require_separate_approver'] is False
    cookie = ok.headers['set-cookie'].lower()
    assert 'httponly' in cookie and 'samesite=strict' in cookie and 'max-age=43200' in cookie
    token = anon_client.cookies.get(settings.SESSION_COOKIE)
    with sync_engine.connect() as conn:
        stored = conn.execute(text('SELECT token_hash FROM user_sessions')).scalar_one()
    assert stored != token and len(stored) == 64  # only the hash is stored
    assert (await anon_client.get('/auth/me')).json()['email'] == 'admin@demo.com'
    assert (await anon_client.post('/clients', json={'name': 'Created after login'})).status_code == 200
    assert (await anon_client.post('/auth/logout')).status_code == 200
    anon_client.cookies.set(settings.SESSION_COOKIE, token)  # a copied token is useless after logout
    assert (await anon_client.get('/clients')).status_code == 401
    # Expired sessions are rejected.
    await anon_client.post('/auth/login', json={'email': 'admin@demo.com', 'password': 'synthetic-test-password'})
    with sync_engine.begin() as conn:
        conn.execute(text("UPDATE user_sessions SET expires_at = NOW() - interval '1 minute' WHERE revoked_at IS NULL"))
    assert (await anon_client.get('/clients')).status_code == 401
    # Five wrong passwords lock the email, even for the right password, and the lock is per email.
    for _ in range(5):
        assert (await anon_client.post('/auth/login', json={'email': 'admin@demo.com', 'password': 'wrong-password'})).status_code == 401
    locked = await anon_client.post('/auth/login', json={'email': 'admin@demo.com', 'password': 'synthetic-test-password'})
    assert locked.status_code == 429 and 'Try again' in locked.json()['detail']
    # Disabled users cannot log in or keep using a session.
    anon_client.cookies.clear()
    _make_user('partner@demo.com', 'another-synthetic-pass')
    await anon_client.post('/auth/login', json={'email': 'partner@demo.com', 'password': 'another-synthetic-pass'})
    assert (await anon_client.get('/clients')).status_code == 200
    with sync_engine.begin() as conn:
        conn.execute(text("UPDATE users SET is_active = false WHERE email = 'partner@demo.com'"))
    assert (await anon_client.get('/clients')).status_code == 401

@pytest.mark.asyncio
async def test_logged_in_firm_sees_only_its_own_data(anon_client):
    _make_user('firm-a@demo.com', 'firm-a-synthetic-pass')
    _make_user('firm-b@demo.com', 'firm-b-synthetic-pass')
    await anon_client.post('/auth/login', json={'email': 'firm-a@demo.com', 'password': 'firm-a-synthetic-pass'})
    created = (await anon_client.post('/clients', json={'name': 'Firm A only'})).json()
    await anon_client.post('/auth/logout'); anon_client.cookies.clear()
    await anon_client.post('/auth/login', json={'email': 'firm-b@demo.com', 'password': 'firm-b-synthetic-pass'})
    assert (await anon_client.get('/clients')).json() == []
    assert (await anon_client.get(f"/clients/{created['id']}/registrations")).status_code == 404


@pytest.mark.asyncio
async def test_setoff_follows_only_ca_confirmed_rules_and_invalidates_approval(client):
    from pathlib import Path
    from io import BytesIO
    from openpyxl import load_workbook
    data = json.loads(Path('tests/fixtures/setoff_v1/cases.json').read_text(encoding='utf-8'))
    expected = data['integration_case']['expected']
    period = await _period(client, 'Setoff client')
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()['id']
    await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    for row in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']:
        await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': 'reviewed' if row['validation_status'] == 'ready' else 'excluded', 'note': 'Synthetic'})
    b = {}
    for kind, path in (('purchase', 'sample_data/v1/purchase_register.csv'), ('statement', 'sample_data/v1/gstr2b_demo.csv')):
        b[kind] = (await client.post(f'/periods/{period}/imports', params={'source_type': kind}, files={'file': (kind, Path(path).read_bytes())})).json()['id']
        await client.post(f'/imports/{b[kind]}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': b['purchase'], 'statement_batch_id': b['statement']})).json()['id']
    results = (await client.get(f'/runs/{run}/itc-decisions')).json()
    items = [{'result_id': r['result_id'], 'decision': 'claim' if r['purchase_record_ids'] == ['PUR-001'] and r['status'] == 'matched' else 'not_claimed'} for r in results]
    assert (await client.post(f'/runs/{run}/itc-decisions', json={'items': items, 'note': 'Synthetic'})).status_code == 200
    adj = f'/periods/{period}/adjustments'
    assert (await client.post(adj, json={'adjustment_type': 'rcm_liability', 'tax_head': 'cgst', 'amount': '5.00', 'note': 'Synthetic RCM'})).status_code == 200
    # Opening credit: user-reported ledger balance; adds to credit.
    assert (await client.post(adj, json={'adjustment_type': 'opening_credit', 'tax_head': 'igst', 'amount': '50.00', 'note': 'Ledger balance brought forward (user-reported)'})).status_code == 200
    params = {'sales_batch_id': sales, 'run_id': run}
    preview = (await client.get(f'/periods/{period}/worksheet', params=params)).json()
    assert preview['payload']['worksheet']['heads']['igst']['opening_credit'] == '50.00' and preview['payload']['worksheet']['heads']['igst']['net'] == '40.00'
    # No order confirmed: set-off is not computed, and the worksheet is still approvable.
    assert preview['payload']['setoff']['status'] == 'not_computed' and 'utilisation order' in preview['payload']['setoff']['reason']
    draft = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    approval = (await client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'Approved before set-off rules'})).json()['id']

    confirm = lambda key, value, **kw: client.post(f'/legal-rules/{key}/confirm', json={'value': value, 'source_reference': 'Test configuration, not law', 'effective_from': '2026-04', 'checked_against_current_law': True} | kw)
    assert (await confirm('credit_utilisation_order', {'steps': ['IGST>UTGST']})).status_code == 400
    assert (await confirm('credit_utilisation_order', {'steps': ['IGST>CGST', 'igst>cgst']})).status_code == 400
    assert (await confirm('payment_rounding', {'multiple': '0', 'direction': 'half_up'})).status_code == 400
    assert (await confirm('payment_rounding', {'multiple': '1.00', 'direction': 'sideways'})).status_code == 400
    assert (await confirm('credit_utilisation_order', {'steps': data['configs']['A']})).status_code == 200
    assert (await confirm('payment_rounding', {'multiple': '1.00', 'direction': 'half_up'})).status_code == 200
    register = {r['key']: r for r in (await client.get('/legal-rules')).json()['rules']}
    assert register['credit_utilisation_order']['active']['value']['steps'][1] == 'IGST>CGST'
    assert register['payment_rounding']['options'] == {'direction': ['half_up', 'up', 'down']}

    # Confirming set-off rules makes the earlier approval out of date; its frozen payload is unchanged.
    drafts = (await client.get(f'/periods/{period}/tax-drafts')).json()
    assert drafts[0]['state'] == 'approved_stale' and drafts[0]['payload']['setoff']['status'] == 'not_computed'
    so = (await client.get(f'/periods/{period}/worksheet', params=params)).json()['payload']['setoff']
    assert so['status'] == 'computed'
    assert [[u['credit_head'], u['liability_head'], u['amount']] for u in so['utilisation']] == expected['utilisation']
    for h, v in expected['cash'].items():
        assert so['heads'][h]['cash'] == v and so['heads'][h]['cash_before_rounding'] == expected['cash_before_rounding'][h] and so['heads'][h]['carry_forward'] == expected['carry_forward'][h]
    assert so['total_cash'] == expected['total_cash'] and so['rules']['credit_utilisation_order']['source_reference'] == 'Test configuration, not law'
    # GSTR-3B view: 3.1(a) = reviewed sales (taxable 1750.00 from the sales README), 3.1(d) = the CGST 5.00 reverse charge,
    # 4(A)(5) = claimed PUR-001, 6.1 = the set-off above. Unsupported rows are not captured (None), never zero.
    g = (await client.get(f'/periods/{period}/worksheet', params=params)).json()['payload']['gstr3b']
    a, d = g['table_3_1'][0], g['table_3_1'][3]
    assert (a['taxable_value'], a['igst'], a['cgst'], a['sgst'], a['cess']) == ('1750.00', '90.00', '112.50', '112.50', '0.00')
    assert d['cgst'] == '5.00' and d['taxable_value'] is None and g['table_3_1'][1]['igst'] is None
    assert g['table_4'][1]['cgst'] == '90.00' and g['table_4'][4]['cgst'] == '90.00' and g['table_4'][0]['igst'] is None
    cgst_pay = next(p for p in g['table_6_1'] if p['head'] == 'cgst')
    assert cgst_pay['tax_payable'] == '117.50' and cgst_pay['paid_through_itc']['cgst'] == '90.00' and cgst_pay['paid_in_cash'] == '28.00' and cgst_pay['interest'] is None

    # A rule version that starts after this period is not in force for it.
    assert (await confirm('credit_utilisation_order', {'steps': data['configs']['B']}, effective_from='2026-09')).status_code == 200
    assert (await client.get(f'/periods/{period}/worksheet', params=params)).json()['payload']['setoff']['status'] == 'not_computed'
    assert (await confirm('credit_utilisation_order', {'steps': data['configs']['A']})).status_code == 200

    # Reopen, redraft, approve; the workbook carries the set-off sheet.
    assert (await client.post(f'/tax-approvals/{approval}/reopen', json={'reason': 'Set-off rules confirmed'})).status_code == 200
    second = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    assert second['payload']['setoff']['total_cash'] == expected['total_cash']
    assert (await client.post(f"/tax-drafts/{second['id']}/approve", json={'note': 'Approved with set-off'})).status_code == 200
    wb = load_workbook(BytesIO((await client.get(f"/tax-drafts/{second['id']}/export")).content))
    sheet = [r for r in wb['Set-off'].iter_rows(min_row=2, values_only=True)]
    assert sheet[-1][0] == 'total' and sheet[-1][3] == expected['total_cash'] and sheet[-1][1] == '90.00'
    assert not any(c.data_type == 'f' for ws in wb for row in ws for c in row)

    # Reversal larger than credit: set-off refuses to guess.
    await client.post(adj, json={'adjustment_type': 'itc_reversal', 'tax_head': 'sgst', 'amount': '200.00', 'note': 'Synthetic reversal'})
    neg = (await client.get(f'/periods/{period}/worksheet', params=params)).json()['payload']['setoff']
    assert neg['status'] == 'not_computed' and 'SGST' in neg['reason']


async def _ready_period(client, reg, code, claim_all=True, decide=True):
    """Period with reviewed sales and decided ITC (all exact matches claimed by default)."""
    from pathlib import Path
    period = (await client.post(f'/registrations/{reg}/periods', json={'period_code': code})).json()['id']
    sales_bytes = Path('sample_data/sales_v1/sales_register.csv').read_bytes().replace(b'2026-08-', f'{code}-'.encode())
    sales = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('s.csv', sales_bytes)})).json()['id']
    await client.post(f'/sales-imports/{sales}/commit', json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    for row in (await client.get(f'/sales-imports/{sales}', params={'limit': 100})).json()['items']:
        await client.post(f"/sales-imports/{sales}/rows/{row['id']}/review", json={'decision': 'reviewed' if row['validation_status'] == 'ready' else 'excluded', 'note': 'Synthetic'})
    b = {}
    for kind, path in (('purchase', 'sample_data/v1/purchase_register.csv'), ('statement', 'sample_data/v1/gstr2b_demo.csv')):
        b[kind] = (await client.post(f'/periods/{period}/imports', params={'source_type': kind}, files={'file': (kind, Path(path).read_bytes())})).json()['id']
        await client.post(f'/imports/{b[kind]}/commit', params={'acknowledge_invalid': True, 'note': 'Synthetic'})
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': b['purchase'], 'statement_batch_id': b['statement']})).json()['id']
    if decide:
        items = [{'result_id': r['result_id'], 'decision': 'claim' if r['status'] == 'matched' and claim_all else 'not_claimed'} for r in (await client.get(f'/runs/{run}/itc-decisions')).json()]
        await client.post(f'/runs/{run}/itc-decisions', json={'items': items, 'note': 'Synthetic'})
    return period, {'sales_batch_id': sales, 'run_id': run}

@pytest.mark.asyncio
async def test_credit_carry_forward_and_due_dates_follow_confirmed_rules(client):
    from datetime import date
    data = json.loads(open('tests/fixtures/setoff_v1/cases.json', encoding='utf-8').read())
    c = (await client.post('/clients', json={'name': 'Carry client'})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin': 'DEMO-CARRY'})).json()['id']
    july, jp = await _ready_period(client, reg, '2026-07')
    aug, ap = await _ready_period(client, reg, '2026-08')
    cf_url = f'/periods/{aug}/carry-forward'
    assert (await client.get(cf_url)).json()['reason'].startswith('2026-07 has no active approval')
    jd = (await client.post(f'/periods/{july}/tax-drafts', json=jp)).json()
    await client.post(f"/tax-drafts/{jd['id']}/approve", json={'note': 'July approved without set-off'})
    nocomp = (await client.get(cf_url)).json()
    assert nocomp['status'] == 'not_available' and 'no computed set-off' in nocomp['reason']
    assert (await client.post(cf_url)).status_code == 400

    confirm = lambda key, value: client.post(f'/legal-rules/{key}/confirm', json={'value': value, 'source_reference': 'Test configuration, not law', 'effective_from': '2026-04', 'checked_against_current_law': True})
    assert (await confirm('credit_utilisation_order', {'steps': data['configs']['A']})).status_code == 200
    # July's approval is now out of date (rules changed) — carrying from it is refused until re-approved.
    assert 'out of date' in (await client.get(cf_url)).json()['reason']
    approval = next(d for d in (await client.get(f'/periods/{july}/tax-drafts')).json() if d['approval'])['approval']['id']
    await client.post(f'/tax-approvals/{approval}/reopen', json={'reason': 'Set-off rules confirmed'})
    jd2 = (await client.post(f'/periods/{july}/tax-drafts', json=jp)).json()
    assert (await client.post(f"/tax-drafts/{jd2['id']}/approve", json={'note': 'July re-approved with set-off'})).status_code == 200
    carry = {h: v['carry_forward'] for h, v in jd2['payload']['setoff']['heads'].items() if float(v['carry_forward']) > 0}
    assert carry  # all exact matches claimed: credit exceeds July liability

    avail = (await client.get(cf_url)).json()
    assert avail['status'] == 'available' and avail['amounts'] == carry and avail['draft_id'] == jd2['id']
    applied = await client.post(cf_url)
    assert applied.status_code == 200 and applied.json()['status'] == 'already_entered'
    assert (await client.post(cf_url)).status_code == 409
    heads = (await client.get(f'/periods/{aug}/worksheet', params=ap)).json()['payload']['worksheet']['heads']
    assert {h: heads[h]['opening_credit'] for h in carry} == carry
    assert all('Brought forward from 2026-07' in a['note'] for a in applied.json()['existing'])

    # Due dates appear on the Board only after a CA confirms them.
    rows = {r['period_id']: r for r in (await client.get('/board')).json()}
    assert rows[aug]['due'] is None
    assert (await confirm('return_due_dates', {'gstr1_day': 31, 'gstr3b_day': 20})).status_code == 200
    rows = {r['period_id']: r for r in (await client.get('/board')).json()}
    due = rows[aug]['due']
    assert due['gstr3b']['date'] == '2026-09-20' and due['gstr1']['date'] == '2026-09-30'  # 31 clamps to September's last day
    assert due['gstr3b']['days_left'] == (date(2026, 9, 20) - date.today()).days
    assert due['source'] == 'Test configuration, not law'


@pytest.mark.asyncio
async def test_roles_separate_approver_and_assignment(client, mock_user_session):
    owner = await mock_user_session()
    added = {}
    for role, name in (('reviewer', 'Riya Reviewer'), ('preparer', 'Pranav Preparer')):
        r = await client.post('/firm/users', json={'email': f'{role}@demo.example', 'display_name': name, 'role': role, 'initial_password': 'synthetic-pass-123'})
        assert r.status_code == 200, r.text
        added[role] = r.json()
    assert (await client.post('/firm/users', json={'email': 'reviewer@demo.example', 'display_name': 'x', 'role': 'reviewer', 'initial_password': 'synthetic-pass-123'})).status_code == 409
    assert (await client.post('/firm/users', json={'email': 'short@demo.example', 'display_name': 'x', 'role': 'reviewer', 'initial_password': 'short'})).status_code == 422
    assert not any('synthetic-pass' in (e['summary'] or '') for e in (await client.get('/audit-events', params={'limit': 200})).json())
    team = (await client.get('/firm')).json()
    assert {u['role'] for u in team['users']} == {'owner', 'reviewer', 'preparer'} and team['require_separate_approver'] is False
    assert (await client.patch(f"/firm/users/{owner.id}", json={'role': 'preparer'})).status_code == 400  # last owner stays

    def as_user(role):
        u = added[role]
        async def current(): return User(id=uuid.UUID(u['id']), email=u['email'], organization_id=owner.organization_id, role=role)
        app.dependency_overrides[get_current_user] = current

    c = (await client.post('/clients', json={'name': 'Roles client'})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin': 'DEMO-ROLES'})).json()['id']
    period, params = await _ready_period(client, reg, '2026-08')
    try:
        # Preparer: drafts and assigns work, but cannot approve, confirm legal values or manage the firm.
        as_user('preparer')
        draft = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
        assert (await client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'x'})).status_code == 403
        assert (await client.post('/legal-rules/return_due_dates/confirm', json={'value': {'gstr1_day': 11, 'gstr3b_day': 20}, 'source_reference': 'Test configuration, not law', 'effective_from': '2026-04', 'checked_against_current_law': True})).status_code == 403
        assert (await client.patch('/firm/settings', json={'require_separate_approver': True})).status_code == 403
        assert (await client.post('/firm/users', json={'email': 'x@demo.example', 'display_name': 'x', 'role': 'owner', 'initial_password': 'synthetic-pass-123'})).status_code == 403
        assert (await client.post(f'/periods/{period}/assign', json={'user_id': added['preparer']['id']})).status_code == 200
        # Owner turns on the separate-approver policy; then drafts and cannot approve their own draft.
        app.dependency_overrides[get_current_user] = mock_user_session
        assert (await client.patch('/firm/settings', json={'require_separate_approver': True})).json()['require_separate_approver'] is True
        own = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
        denied = await client.post(f"/tax-drafts/{own['id']}/approve", json={'note': 'Own work'})
        assert denied.status_code == 403 and 'Firm policy' in denied.json()['detail']
        # A reviewer who did not draft it may approve.
        as_user('reviewer')
        assert (await client.post(f"/tax-drafts/{own['id']}/approve", json={'note': 'Reviewed by a second person'})).status_code == 200
        approval = next(d for d in (await client.get(f'/periods/{period}/tax-drafts')).json() if d['approval'])['approval']
        as_user('preparer')
        assert (await client.post(f"/tax-approvals/{approval['id']}/reopen", json={'reason': 'x'})).status_code == 403
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session
    row = next(r for r in (await client.get('/board')).json() if r['period_id'] == period)
    assert row['assignee']['name'] == 'Pranav Preparer'
    # Disabling a user ends their access; unassigning works.
    assert (await client.patch(f"/firm/users/{added['preparer']['id']}", json={'is_active': False})).json()['is_active'] is False
    assert (await client.post(f'/periods/{period}/assign', json={'user_id': added['preparer']['id']})).status_code == 404
    assert (await client.post(f'/periods/{period}/assign', json={'user_id': None})).json()['assignee'] is None
    events = {e['action'] for e in (await client.get('/audit-events', params={'limit': 200})).json()}
    assert {'user_added', 'firm_policy_changed', 'period_assigned', 'user_changed', 'tax_draft_approved'} <= events


@pytest.mark.asyncio
async def test_ims_actions_gate_itc_claims_suggestions_and_approval(client, mock_user_session):
    c = (await client.post('/clients', json={'name': 'IMS client'})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin': 'DEMO-IMS'})).json()['id']
    period, params = await _ready_period(client, reg, '2026-08', decide=False)
    run = (await client.get(f"/periods/{period}/reconciliation-runs/{params['run_id']}")).json()
    inbox_url = f"/imports/{run['statement_batch_id']}/ims"
    assert (await client.get(f"/imports/{run['purchase_batch_id']}/ims")).status_code == 400  # purchases are not IMS invoices
    inbox = (await client.get(inbox_url)).json()
    assert len(inbox['items']) == 19 and inbox['counts']['no_action'] == 19  # 20 statement rows, 1 invalid
    by_rid = {i['record_id']: i for i in inbox['items']}
    assert by_rid['STMT-001']['finding'] == 'matched' and by_rid['STMT-010']['finding'] == 'amount_mismatch'

    # Pending and rejected invoices change the pre-fill: no claim for them.
    act = lambda items, note='Synthetic IMS review': client.post(inbox_url, json={'items': items, 'note': note})
    assert (await act([{'record_id': by_rid['STMT-001']['id'], 'action': 'reject'}, {'record_id': by_rid['STMT-002']['id'], 'action': 'pending'}])).status_code == 200
    assert (await act([{'record_id': by_rid['STMT-001']['id'], 'action': 'accept'}])).status_code == 409  # stale previous id
    sugg = {s['result_id']: s for s in (await client.get(f"/runs/{params['run_id']}/itc-suggestions")).json()['suggestions']}
    results = (await client.get(f"/runs/{params['run_id']}/itc-decisions")).json()
    r1 = next(r for r in results if r['statement_record_ids'] == ['STMT-001']); r2 = next(r for r in results if r['statement_record_ids'] == ['STMT-002'])
    assert sugg[r1['result_id']]['decision'] == 'not_claimed' and 'rejected in IMS' in sugg[r1['result_id']]['reason']
    assert sugg[r2['result_id']]['decision'] == 'deferred'
    # Claiming a rejected invoice is refused.
    bad = await client.post(f"/runs/{params['run_id']}/itc-decisions", json={'items': [{'result_id': r1['result_id'], 'decision': 'claim'}], 'note': 'x'})
    assert bad.status_code == 400 and 'rejected in IMS' in bad.json()['detail']

    # Accept STMT-001 again, claim exact matches; then a later IMS reject blocks and makes the approval out of date.
    latest1 = (await client.get(inbox_url)).json()['items']
    s1 = next(i for i in latest1 if i['record_id'] == 'STMT-001'); s2 = next(i for i in latest1 if i['record_id'] == 'STMT-002')
    assert (await act([{'record_id': s1['id'], 'action': 'accept', 'previous_action_id': s1['previous_action_id']}, {'record_id': s2['id'], 'action': 'accept', 'previous_action_id': s2['previous_action_id']}])).status_code == 200
    items = [{'result_id': r['result_id'], 'decision': 'claim' if r['status'] == 'matched' else 'not_claimed'} for r in results]
    assert (await client.post(f"/runs/{params['run_id']}/itc-decisions", json={'items': items, 'note': 'Synthetic'})).status_code == 200
    draft = (await client.post(f'/periods/{period}/tax-drafts', json=params)).json()
    assert draft['blockers'] == []
    assert (await client.post(f"/tax-drafts/{draft['id']}/approve", json={'note': 'Approved'})).status_code == 200
    s1 = next(i for i in (await client.get(inbox_url)).json()['items'] if i['record_id'] == 'STMT-001')
    assert (await act([{'record_id': s1['id'], 'action': 'reject', 'previous_action_id': s1['previous_action_id']}], 'Supplier invoice disputed')).status_code == 200
    assert (await client.get(f'/periods/{period}/tax-drafts')).json()[0]['state'] == 'approved_stale'
    blockers = (await client.get(f'/periods/{period}/worksheet', params=params)).json()['blockers']
    assert any('rejected or pending in IMS' in b and 'STMT-001' in b for b in blockers)
    assert 'ims_action' in {e['action'] for e in (await client.get('/audit-events', params={'limit': 200})).json()}

    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4(), role='owner')
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.get(inbox_url)).status_code == 404
        assert (await act([{'record_id': s1['id'], 'action': 'accept'}])).status_code == 404
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session

@pytest.mark.asyncio
async def test_gstr2b_json_import_reconciles_against_gstin_purchase_register(client, mock_user_session):
    from pathlib import Path
    period = await _period(client, 'GSTR-2B client')
    sample = Path('tests/fixtures/gstr2b_v1/sample_gstr2b.json').read_bytes()
    r = await client.post(f'/periods/{period}/imports/gstr2b', files={'file': ('GSTR2B_082026.json', sample)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body['batch']['record_count'] == 3 and body['batch']['invalid_count'] == 0 and body['conversion']['converted'] == 3
    assert len(body['conversion']['skipped']) == 3 and body['conversion']['itc_unavailable'][0]['record_id'] == '2B-0002'
    statement = body['batch']['id']
    assert (await client.post(f'/imports/{statement}/commit')).status_code == 200
    # Purchase register keyed by the same GSTINs: two match exactly, one is books-only.
    purchases = ('record_id,document_type,supplier_ref,invoice_number,invoice_date,taxable_value,cgst,sgst,igst,cess,invoice_total,description\n'
                 'P-1,invoice,00BBBBB1111B1Z1,SUP1/001,2026-08-03,2000.00,180.00,180.00,0.00,0.00,2360.00,Stationery\n'
                 'P-2,invoice,00CCCCC2222C1Z2,0043,2026-08-15,1000.00,0.00,0.00,120.00,1.00,1121.00,Freight\n'
                 'P-3,invoice,00CCCCC2222C1Z2,0044,2026-08-20,500.00,45.00,45.00,0.00,0.00,590.00,Not filed by supplier\n').encode()
    p = (await client.post(f'/periods/{period}/imports', params={'source_type': 'purchase'}, files={'file': ('p.csv', purchases)})).json()
    assert p['invalid_count'] == 0  # GSTIN-shaped supplier references validate
    await client.post(f"/imports/{p['id']}/commit")
    run = (await client.post(f'/periods/{period}/reconciliation-runs', json={'purchase_batch_id': p['id'], 'statement_batch_id': statement})).json()
    counts = run['summary_data']['distinct_results_by_status']
    assert counts.get('matched') == 2 and counts.get('books_only') == 1 and counts.get('statement_only') == 1
    assert 'gstr2b_converted' in {e['action'] for e in (await client.get('/audit-events', params={'limit': 200})).json()}
    # Wrong month and non-2B files are refused with a reason.
    bad = await client.post(f'/periods/{period}/imports/gstr2b', files={'file': ('x.json', b'{"a": 1}')})
    assert bad.status_code == 400 and 'docdata' in bad.json()['detail']
    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4(), role='owner')
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.post(f'/periods/{period}/imports/gstr2b', files={'file': ('x.json', sample)})).status_code == 404
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session


@pytest.mark.asyncio
async def test_sales_v2_gstr1_draft_export(client, mock_user_session):
    from pathlib import Path
    expected = json.loads(Path('tests/fixtures/sales_v2/expected_gstr1.json').read_text(encoding='utf-8'))
    c = (await client.post('/clients', json={'name': 'GSTR-1 client'})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin': expected['gstin']})).json()['id']
    period = (await client.post(f'/registrations/{reg}/periods', json={'period_code': '2026-08'})).json()['id']
    template = await client.get('/sales/template', params={'version': '2'})
    assert template.text.strip().endswith('customer_gstin,pos_state_code,rate,hsn_code,uqc,quantity')
    v1 = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('v1.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})).json()
    assert v1['contract_version'] == 'sales-v1'
    await client.post(f"/sales-imports/{v1['id']}/commit", json={'acknowledge_blocked': True, 'note': 'Synthetic'})
    refused = await client.get(f"/sales-imports/{v1['id']}/gstr1")
    assert refused.status_code == 400 and 'v2 template' in refused.json()['detail']
    b = (await client.post(f'/periods/{period}/sales-imports', files={'file': ('v2.csv', Path('sample_data/sales_v2/sales_register.csv').read_bytes())})).json()
    assert b['contract_version'] == 'sales-v2'
    url = f"/sales-imports/{b['id']}"
    await client.post(url + '/commit', json={'acknowledge_blocked': True, 'note': 'Two rows with bad GSTIN data'})
    items = (await client.get(url, params={'limit': 100})).json()['items']
    for row in items:
        if row['raw_data']['record_id'] == 'S2-006': continue  # left pending: must not be included
        await client.post(f"{url}/rows/{row['id']}/review", json={'decision': 'reviewed' if row['validation_status'] == 'ready' else 'excluded', 'note': 'Synthetic'})
    out = (await client.get(url + '/gstr1')).json()
    assert out['included_rows'] == 5 and any('1 row(s) still pending' in w for w in out['warnings'])
    assert out['document']['b2b'] == expected['b2b'] and out['document']['fp'] == '082026' and out['document']['gstin'] == expected['gstin']
    assert all(g['pos'] != '01' for g in out['document']['b2cs'])  # S2-006 (inter-state B2C) is pending, so not exported
    row6 = next(r for r in items if r['raw_data']['record_id'] == 'S2-006')
    await client.post(f"{url}/rows/{row6['id']}/review", json={'decision': 'reviewed', 'note': 'Synthetic'})
    full = (await client.get(url + '/gstr1')).json()
    assert full['document']['b2cs'] == expected['b2cs'] and full['document']['hsn'] == expected['hsn']
    assert 'gstr1_draft_exported' in {e['action'] for e in (await client.get('/audit-events', params={'limit': 200})).json()}


@pytest.mark.asyncio
async def test_client_upload_links_and_reminder_drafts(client, mock_user_session):
    import re
    from pathlib import Path
    from datetime import datetime, timedelta, timezone
    c = (await client.post('/clients', json={'name': 'Links client'})).json()['id']
    reg = (await client.post(f'/clients/{c}/registrations', json={'gstin': 'DEMO-LINKS'})).json()['id']
    period = (await client.post(f'/registrations/{reg}/periods', json={'period_code': '2026-08'})).json()['id']
    assert (await client.patch(f'/clients/{c}/contact', json={'contact_email': 'not-an-email'})).status_code == 422
    assert (await client.patch(f'/clients/{c}/contact', json={'contact_email': 'owner@client.example', 'contact_phone': '+91 98765 43210'})).status_code == 200

    draft = (await client.post(f'/periods/{period}/reminder', json={'base_url': 'http://127.0.0.1:3000'})).json()
    assert draft['missing'] == ['sales', 'purchase'] and 'August 2026' in draft['subject']
    tokens = re.findall(r'/u/([A-Za-z0-9_-]+)', draft['body'])
    assert len(tokens) == 2 and draft['mailto'].startswith('mailto:owner%40client.example?subject=') and draft['whatsapp'].startswith('https://wa.me/919876543210?text=')
    assert 'does not send' in draft['note']
    assert not any(t in (e['summary'] or '') for e in (await client.get('/audit-events', params={'limit': 200})).json() for t in tokens)  # tokens never logged

    sales_token, purchase_token = tokens
    info = (await client.get(f'/public/upload/{sales_token}')).json()
    assert info['client'] == 'Links client' and info['kind'] == 'sales' and info['uses_left'] == 5
    up = await client.post(f'/public/upload/{sales_token}', files={'file': ('sales.csv', Path('sample_data/sales_v1/sales_register.csv').read_bytes())})
    assert up.status_code == 200 and up.json()['status'] == 'received'
    batches = (await client.get(f'/periods/{period}/sales-imports')).json()
    assert len(batches) == 1 and batches[0]['status'] == 'preview'  # staff still review and commit
    assert (await client.post(f'/public/upload/{purchase_token}', files={'file': ('p.csv', Path('sample_data/v1/purchase_register.csv').read_bytes())})).status_code == 200
    assert [b['status'] for b in (await client.get(f'/periods/{period}/imports')).json()] == ['preview']
    summaries = [e['summary'] or '' for e in (await client.get('/audit-events', params={'period_id': period, 'limit': 200})).json()]
    assert sum('uploaded by the client via link' in s for s in summaries) == 2
    # Nothing is awaiting files any more, so a new reminder has nothing to ask for.
    assert (await client.post(f'/periods/{period}/reminder', json={})).json()['missing'] == []

    one = (await client.post(f'/periods/{period}/upload-links', json={'kind': 'sales', 'max_uses': 1})).json()
    tok = one['path'].split('/u/')[1]
    bad = await client.post(f'/public/upload/{tok}', files={'file': ('x.csv', b'wrong,header\n1,2')})
    assert bad.status_code == 400  # a rejected file does not use up the link
    assert (await client.get(f'/public/upload/{tok}')).json()['uses_left'] == 1
    v2 = Path('sample_data/sales_v1/sales_register.csv').read_bytes().replace(b'S-001', b'LINK-001', 1)
    assert (await client.post(f'/public/upload/{tok}', files={'file': ('s.csv', v2)})).status_code == 200
    assert (await client.get(f'/public/upload/{tok}')).status_code == 410
    links = {l['id']: l for l in (await client.get(f'/periods/{period}/upload-links')).json()}
    assert links[one['id']]['state'] == 'used_up' and 'token' not in str(links)
    another = (await client.post(f'/periods/{period}/upload-links', json={'kind': 'purchase'})).json()
    assert (await client.post(f"/upload-links/{another['id']}/revoke")).json()['state'] == 'revoked'
    assert (await client.get(f"/public/upload/{another['path'].split('/u/')[1]}")).status_code == 410
    later = (await client.post(f'/periods/{period}/upload-links', json={'kind': 'purchase'})).json()
    with sync_engine.begin() as conn:
        conn.execute(text('UPDATE upload_links SET expires_at = :t WHERE id = :i'), {'t': datetime.now(timezone.utc) - timedelta(minutes=1), 'i': uuid.UUID(later['id'])})
    assert (await client.get(f"/public/upload/{later['path'].split('/u/')[1]}")).status_code == 410
    assert (await client.get('/public/upload/not-a-real-token')).status_code == 404

    user = await mock_user_session()
    async def foreign_user(): return User(id=user.id, email=user.email, organization_id=uuid.uuid4(), role='owner')
    app.dependency_overrides[get_current_user] = foreign_user
    try:
        assert (await client.post(f'/periods/{period}/upload-links', json={'kind': 'sales'})).status_code == 404
        assert (await client.post(f'/periods/{period}/reminder', json={})).status_code == 404
        assert (await client.post(f"/upload-links/{one['id']}/revoke")).status_code == 404
    finally:
        app.dependency_overrides[get_current_user] = mock_user_session
