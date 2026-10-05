"""Excel (.xlsx) uploads: the first sheet is converted to the same CSV the CSV contracts expect, then goes through
the normal parsing and validation. Conversion never fixes data: amounts with more than two decimals or text in
number columns pass through unchanged so validation rejects them visibly."""
import csv
import hashlib
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS = 20000
MAX_COLS = 100
# Money columns of the purchase, statement and sales contracts. Excel drops trailing zeros (1000.00 is stored as
# 1000), so these get two decimals back; a value with more than two decimals is left as is for validation to reject.
MONEY = {'taxable_value', 'cgst', 'sgst', 'igst', 'cess', 'invoice_total'}


def is_xlsx(content: bytes) -> bool:
    return content[:4] == b'PK\x03\x04'


def _cell(value) -> str:
    if value is None: return ''
    if isinstance(value, bool): return 'TRUE' if value else 'FALSE'
    if isinstance(value, int): return str(value)
    if isinstance(value, float):
        # repr() is the shortest text that round-trips, so 1234.56 stays 1234.56 rather than 1234.5599999...
        return str(int(value)) if value.is_integer() else format(Decimal(repr(value)), 'f')
    if isinstance(value, datetime): return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, (date, time)): return value.isoformat()
    return str(value)


def _money(value) -> str:
    text = _cell(value)
    if isinstance(value, bool) or not isinstance(value, (int, float)): return text
    amount = Decimal(text)
    return f'{amount:.2f}' if amount == amount.quantize(Decimal('0.01')) else text


def to_csv(content: bytes) -> bytes:
    try: book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc: raise ValueError('Could not read the Excel file. Save it as .xlsx (Excel Workbook) or upload CSV.') from exc
    try:
        if not book.worksheets: raise ValueError('The Excel file has no sheets')
        out = StringIO(); writer = csv.writer(out, lineterminator='\n'); width = None; count = 0; money = set()
        for row in book.worksheets[0].iter_rows(values_only=True):
            cells = [_money(v) if i in money else _cell(v) for i, v in enumerate(row[:MAX_COLS])]
            if not any(c.strip() for c in cells): continue
            if width is None:  # header row
                width = max(i + 1 for i, c in enumerate(cells) if c.strip())
                money = {i for i, c in enumerate(cells) if c.strip().lower() in MONEY}
            count += 1
            if count > MAX_ROWS + 1: raise ValueError(f'The Excel sheet has more than {MAX_ROWS} rows')
            writer.writerow(cells[:width] + [''] * (width - len(cells)))
        if width is None: raise ValueError('The first sheet of the Excel file is empty')
        return out.getvalue().encode('utf-8')
    finally:
        book.close()


def convert_upload(filename: str | None, content: bytes) -> tuple[str, bytes, bytes | None]:
    """Returns (filename, csv bytes, original workbook or None). Non-Excel content passes through. The batch
    filename records the workbook hash; call keep_original once the import is accepted."""
    if not is_xlsx(content): return filename or '', content, None
    csv_bytes = to_csv(content)
    name = (filename or 'upload.xlsx').replace('\\', '/').split('/')[-1][:120]
    return f'{name} (converted from Excel, sha256 {hashlib.sha256(content).hexdigest()[:16]})', csv_bytes, content


def keep_original(original: bytes | None, storage_dir: str, org_hex: str) -> None:
    if original is None: return
    path = Path(storage_dir) / org_hex / f'{hashlib.sha256(original).hexdigest()}.xlsx'
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists(): path.write_bytes(original)