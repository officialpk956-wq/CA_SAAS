"""Organization-scoped sales preparation; independent of purchase matching."""
import csv
import hashlib
from collections import defaultdict
from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import Font
from ..dependencies import get_db, get_active_organization, get_current_user
from ...db.models import Client, FilingPeriod, GSTRegistration, SalesBatch, SalesRecord, SalesReview, AuditEvent, User
from ...sales import HEADERS, HEADERS_V2, AMOUNTS, MAX_BYTES, VERSION, contract_version, parse_sales, summarize
from ...services.export_service import sanitize_value
from ...config import settings
from ...xlsx import convert_upload, keep_original

router = APIRouter(tags=['sales'])

async def owned_period(db, period_id, org_id, lock=False):
    query = select(FilingPeriod).join(GSTRegistration).join(Client).where(FilingPeriod.id == period_id,Client.organization_id == org_id)
    if lock: query = query.with_for_update(of=FilingPeriod)
    period = (await db.execute(query)).scalars().first()
    if not period: raise HTTPException(404,'Period not found')
    return period

async def owned_batch(db, batch_id, org_id, lock=False):
    query = select(SalesBatch).join(FilingPeriod).join(GSTRegistration).join(Client).where(SalesBatch.id == batch_id,Client.organization_id == org_id)
    if lock: query = query.with_for_update(of=SalesBatch)
    batch = (await db.execute(query)).scalars().first()
    if not batch: raise HTTPException(404,'Sales import not found')
    return batch

def batch_view(batch):
    return {key:str(getattr(batch,key)) for key in ('id','period_id','user_id','filename','file_hash','contract_version','status','created_at')} | {'commit_note':batch.commit_note}

def review_view(review):
    return {'id':str(review.id),'actor_id':str(review.user_id),'decision':review.decision,'note':review.note,'created_at':review.created_at.isoformat()}

async def record_views(db, batch_id):
    records = (await db.execute(select(SalesRecord).where(SalesRecord.batch_id == batch_id).order_by(SalesRecord.row_number))).scalars().all()
    reviews = (await db.execute(select(SalesReview).join(SalesRecord).where(SalesRecord.batch_id == batch_id).order_by(SalesReview.created_at,SalesReview.id))).scalars().all()
    history = defaultdict(list)
    for review in reviews: history[review.record_id].append(review_view(review))
    return [{'id':str(r.id),'row_number':r.row_number,'raw_data':r.raw_data,'validation_status':r.validation_status,'issues':r.issues,'history':history[r.id],'decision':history[r.id][-1]['decision'] if history[r.id] else 'unresolved','previous_review_id':history[r.id][-1]['id'] if history[r.id] else None} for r in records]

@router.get('/sales/template')
async def template(version: Literal['1', '2'] = '1', org_id: UUID = Depends(get_active_organization)):
    stream = StringIO();csv.writer(stream).writerow(HEADERS_V2 if version == '2' else HEADERS)
    return Response(stream.getvalue(),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="synthetic_sales_template.csv"'})

@router.post('/periods/{period_id}/sales-imports')
async def upload(period_id: UUID, file: UploadFile = File(...), org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db,period_id,org_id,lock=True)
    content = await file.read(MAX_BYTES + 1)
    return await store_sales_upload(db, period, org_id, user.id, file.filename, content)

async def store_sales_upload(db, period, org_id, user_id, filename, content: bytes, via: str = ''):
    """Shared by staff uploads and client upload links. The caller holds the period lock."""
    if len(content) > MAX_BYTES: raise HTTPException(413,'Sales file exceeds 5 MiB')
    try: filename, content, workbook = convert_upload(filename, content)
    except ValueError as exc: raise HTTPException(400,str(exc)) from exc
    digest = hashlib.sha256(content).hexdigest()
    existing = (await db.execute(select(SalesBatch).where(SalesBatch.period_id == period.id,SalesBatch.file_hash == digest))).scalars().first()
    if existing: return batch_view(existing)
    try: records = parse_sales(content,period.period_code)
    except ValueError as exc: raise HTTPException(400,str(exc)) from exc
    name = (filename or 'sales.csv').replace('\\','/').split('/')[-1][:200]
    # Keep exact decoded UTF-8 including BOM; re-encoding reproduces the original hash.
    batch = SalesBatch(period_id=period.id,user_id=user_id,filename=name,file_hash=digest,original_csv=content.decode('utf-8'),contract_version=contract_version(content),status='preview')
    db.add(batch);await db.flush()
    db.add_all([SalesRecord(batch_id=batch.id,**record) for record in records])
    db.add(AuditEvent(organization_id=org_id,user_id=user_id,action='sales_uploaded',resource_type='sales_batch',resource_id=str(batch.id),period_id=period.id,summary=f'{len(records)} sales rows{via}'))
    await db.commit();await db.refresh(batch)
    keep_original(workbook, settings.STORAGE_DIR, org_id.hex)
    return batch_view(batch)

@router.get('/periods/{period_id}/sales-imports')
async def batches(period_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    await owned_period(db,period_id,org_id)
    records = (await db.execute(select(SalesBatch).where(SalesBatch.period_id == period_id).order_by(SalesBatch.created_at.desc(),SalesBatch.id.desc()))).scalars().all()
    return [batch_view(b) for b in records]

@router.get('/sales-imports/{batch_id}')
async def detail(batch_id: UUID, offset: int = Query(0,ge=0), limit: int = Query(25,ge=1,le=100), status: Literal['all','ready','invalid','duplicate','unsupported'] = 'all', org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    batch = await owned_batch(db,batch_id,org_id)
    rows = await record_views(db,batch.id)
    filtered = [r for r in rows if status == 'all' or r['validation_status'] == status]
    return {'batch':batch_view(batch),'summary':summarize(rows),'total':len(filtered),'items':filtered[offset:offset+limit]}

class CommitInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    acknowledge_blocked: bool = False
    note: str = Field(default='',max_length=2000)

@router.post('/sales-imports/{batch_id}/commit')
async def commit(batch_id: UUID, data: CommitInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    batch = await owned_batch(db,batch_id,org_id,lock=True)
    if batch.status == 'committed': return batch_view(batch)
    blocked = (await db.execute(select(SalesRecord.id).where(SalesRecord.batch_id == batch.id,SalesRecord.validation_status != 'ready').limit(1))).first()
    if blocked and (not data.acknowledge_blocked or not data.note.strip()): raise HTTPException(400,'Acknowledge blocked rows and provide a note before committing')
    batch.status='committed';batch.commit_note=data.note.strip() or None
    db.add(AuditEvent(organization_id=org_id,user_id=user.id,action='sales_committed',resource_type='sales_batch',resource_id=str(batch.id),period_id=batch.period_id))
    await db.commit();await db.refresh(batch)
    return batch_view(batch)

class ReviewInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    decision: Literal['reviewed','excluded','unresolved']
    note: str = Field(min_length=1,max_length=2000)
    previous_review_id: UUID | None = None

@router.post('/sales-imports/{batch_id}/rows/{record_id}/review')
async def review(batch_id: UUID, record_id: UUID, data: ReviewInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    batch = await owned_batch(db,batch_id,org_id,lock=True)
    if batch.status != 'committed': raise HTTPException(400,'Commit the sales import before reviewing')
    record = (await db.execute(select(SalesRecord).where(SalesRecord.id == record_id,SalesRecord.batch_id == batch.id))).scalars().first()
    if not record: raise HTTPException(404,'Sales row not found')
    if not data.note.strip(): raise HTTPException(400,'A review note is required')
    if data.decision == 'reviewed' and record.validation_status != 'ready': raise HTTPException(400,'Blocked rows cannot be included; correct and upload a new version')
    previous = (await db.execute(select(SalesReview).where(SalesReview.record_id == record.id).order_by(SalesReview.created_at.desc(),SalesReview.id.desc()))).scalars().first()
    if (previous.id if previous else None) != data.previous_review_id: raise HTTPException(409,'A newer review exists. Reload this import before saving.')
    event = SalesReview(record_id=record.id,user_id=user.id,decision=data.decision,note=data.note.strip())
    db.add(event);await db.flush()
    db.add(AuditEvent(organization_id=org_id,user_id=user.id,action='sales_'+data.decision,resource_type='sales_review',resource_id=str(event.id),period_id=batch.period_id,summary=f'row {record.row_number}'))
    await db.commit();await db.refresh(event)
    return review_view(event)

@router.get('/sales-imports/{batch_id}/export')
async def export(batch_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    # Reviews lock the same batch, giving this working paper a consistent cutoff.
    batch = await owned_batch(db,batch_id,org_id,lock=True)
    if batch.status != 'committed': raise HTTPException(400,'Commit the sales import before exporting')
    period = await owned_period(db,batch.period_id,org_id)
    cutoff = datetime.now(timezone.utc).isoformat()
    rows = await record_views(db,batch.id);summary = summarize(rows)
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,headers,values):
        ws=wb.create_sheet(name);ws.append(headers)
        for row in values:
            ws.append([sanitize_value(v) for v in row])
            # Force literal strings even if openpyxl could recognize a formula.
            for cell in ws[ws.max_row]: cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True)
    metadata = {'label':'SYNTHETIC DRAFT - NOT FOR FILING','batch_id':str(batch.id),'period_id':str(period.id),'period_code':period.period_code,'registration_id':str(period.registration_id),'source_hash':batch.file_hash,'source_filename':batch.filename,'contract_version':batch.contract_version,'uploader_id':str(batch.user_id),'review_cutoff':cutoff,'commit_note':batch.commit_note,'scope':'Included totals are reviewed ready source amounts only. No statutory tax determination.'}
    sheet('Metadata',['key','value'],metadata.items())
    sheet('Summary',['measure','count'],[(k,v) for k,v in summary.items() if k != 'included_totals'])
    sheet('Included Totals',['field','amount'],summary['included_totals'].items())
    sheet('Sales Rows',['row_id','row_number',*HEADERS,'validation_status','issues','decision'],[[r['id'],r['row_number'],*[r['raw_data'][k] for k in HEADERS],r['validation_status'],r['issues'],r['decision']] for r in rows])
    sheet('Review History',['row_id','review_id','actor_id','decision','note','created_at'],[[r['id'],h['id'],h['actor_id'],h['decision'],h['note'],h['created_at']] for r in rows for h in r['history']])
    stream=BytesIO();wb.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="Synthetic_Sales_{batch.id}.xlsx"'})


@router.get('/sales-imports/{batch_id}/gstr1')
async def gstr1_draft(batch_id: UUID, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    """GSTR-1 draft JSON from reviewed, ready rows of a committed sales-v2 import. Not validated against the portal."""
    from ...gstr1 import build
    from ...sales import GSTIN_SHAPE
    batch = await owned_batch(db, batch_id, org_id)
    if batch.status != 'committed': raise HTTPException(400, 'Commit the sales import first')
    if batch.contract_version != 'sales-v2': raise HTTPException(400, 'GSTR-1 needs the sales v2 template (customer GSTIN, state code, rate, HSN, unit, quantity). Download it from Sales and re-upload.')
    period = await owned_period(db, batch.period_id, org_id)
    reg = await db.get(GSTRegistration, period.registration_id)
    rows = await record_views(db, batch.id)
    ready = [r['raw_data'] for r in rows if r['decision'] == 'reviewed' and r['validation_status'] == 'ready']
    if not ready: raise HTTPException(400, 'No reviewed, ready rows to include')
    doc, warnings = build(ready, reg.gstin, period.period_code)
    pending = sum(1 for r in rows if r['decision'] == 'unresolved')
    if pending: warnings.insert(0, f'{pending} row(s) still pending review are not included.')
    if not GSTIN_SHAPE.fullmatch(reg.gstin): warnings.insert(0, f'Registration reference {reg.gstin} is not a GSTIN; the portal will reject this file.')
    db.add(AuditEvent(organization_id=org_id, user_id=user.id, action='gstr1_draft_exported', resource_type='sales_batch', resource_id=str(batch.id), period_id=period.id,
                      summary=f"{len(ready)} row(s): {len(doc['b2b'])} B2B customer(s), {len(doc.get('cdnr', []))} CDNR customer(s), {len(doc['b2cs'])} B2CS group(s), {len(doc['hsn']['data'])} HSN line(s)"))
    await db.commit()
    return {'document': doc, 'warnings': warnings, 'included_rows': len(ready),
            'notice': 'Draft built from GSTN\'s published GSTR-1 field names; not validated against the offline tool or portal. Open it in the GSTN offline tool and have a CA review it. Not a filed return.'}
