"""Phase 5A draft tax worksheet: ITC decisions, adjustments, frozen drafts, approval, reopening, filing evidence.

Figures come only from calc-v1 over reviewed inputs. Approval freezes a draft; any later input change makes it stale.
"""
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from ...rules import rule_applies
from ..dependencies import APPROVERS, get_db, get_active_organization, get_current_user, require_role
from .sales import owned_period, record_views as sales_record_views
from ...calculation import ADJUSTMENT_TYPES, ENGINE_VERSION, HEADS, compute, sign
from ...setoff import VERSION as SETOFF_VERSION, parse_steps, setoff
from ...gstr3b import summary as gstr3b_summary
from ...db.models import (AuditEvent, Adjustment, AdjustmentVoid, Client, FilingEvidence, FilingPeriod, GSTRegistration, ImportRecord,
                          ItcDecision, KnowledgeRule, LegalRule, Organization, ReconciliationResult, ReconciliationRun, RuleAcknowledgement, SalesBatch, SalesReview,
                          SalesRecord, TaxApproval, TaxDraft, TaxReopen, User)
from ...services.audit import record as audit
from ...services.ims import blocking as ims_blocking, ims_state
from ...services.export_service import sanitize_value
from ...validation import parse_decimal

router = APIRouter(tags=['worksheet'])
ITC_DECISIONS = ('claim', 'not_claimed', 'deferred')
NOTICE = ('SYNTHETIC DRAFT - NOT FOR FILING. calc-v1 sums reviewed, supplied amounts per tax head; set-off and rounding follow only the '
          'firm\'s CA-confirmed rules (otherwise not computed). No rate or eligibility determination.')
SETOFF_KEYS = ('credit_utilisation_order', 'payment_rounding')

async def setoff_rules(db, org_id, period_code):
    """Active CA-confirmed set-off rules in force for this period, by key."""
    rows = (await db.execute(select(LegalRule).where(LegalRule.organization_id == org_id, LegalRule.status == 'active',
                                                     LegalRule.key.in_(SETOFF_KEYS), LegalRule.effective_from <= period_code))).scalars().all()
    return {r.key: r for r in rows}

def setoff_section(worksheet, found):
    """Set-off result for the payload, or the reason it was not computed. Never guesses an order."""
    if not worksheet: return {'version': SETOFF_VERSION, 'status': 'not_computed', 'reason': 'The worksheet has a calculation error.'}
    order = found.get('credit_utilisation_order')
    if not order:
        return {'version': SETOFF_VERSION, 'status': 'not_computed',
                'reason': 'No CA-confirmed credit utilisation order is in force for this period. Confirm it under Knowledge → Legal rules.'}
    rounding = found.get('payment_rounding')
    heads = worksheet['heads']
    out = setoff({h: heads[h]['liability'] for h in heads}, {h: heads[h]['credit'] for h in heads},
                 parse_steps(order.value['steps']), rounding.value if rounding else None)
    out['rules'] = {k: {'value': r.value, 'source_reference': r.source_reference, 'effective_from': r.effective_from, 'rule_id': str(r.id)} for k, r in found.items()}
    if not rounding: out['rounding_note'] = 'No CA-confirmed rounding rule: cash shown to the paisa.'
    return out

def _decision_view(d):
    return {'id': str(d.id), 'decision': d.decision, 'note': d.note, 'actor_id': str(d.user_id), 'created_at': d.created_at.isoformat()}

async def _owned_run(db, run_id, org_id, lock=False):
    query = select(ReconciliationRun).join(FilingPeriod).join(GSTRegistration).join(Client).where(ReconciliationRun.id == run_id, Client.organization_id == org_id)
    if lock: query = query.with_for_update(of=ReconciliationRun)
    run = (await db.execute(query)).scalars().first()
    if not run: raise HTTPException(404, 'Reconciliation run not found')
    return run

async def _itc_state(db, run):
    results = (await db.execute(select(ReconciliationResult).where(ReconciliationResult.run_id == run.id).order_by(ReconciliationResult.result_id))).scalars().all()
    decisions = (await db.execute(select(ItcDecision).where(ItcDecision.run_id == run.id).order_by(ItcDecision.created_at, ItcDecision.id))).scalars().all()
    history = defaultdict(list)
    for d in decisions: history[d.result_id].append(d)
    return results, history

def _ids(value):
    return value.get('ids', []) if isinstance(value, dict) else value

async def reminder_state(db, period, org_id):
    """Active reminder rules that apply to this period (client-specific or firm-wide), and acknowledgements by rule."""
    client_id = (await db.execute(select(GSTRegistration.client_id).where(GSTRegistration.id == period.registration_id))).scalar_one()
    candidates = (await db.execute(select(KnowledgeRule).where(KnowledgeRule.organization_id == org_id, KnowledgeRule.status == 'active', KnowledgeRule.rule_kind == 'reminder',
                                                            or_(KnowledgeRule.client_id == client_id, KnowledgeRule.client_id.is_(None))))).scalars().all()
    reminders = [r for r in candidates if rule_applies(r.frequency, r.effective_from, r.effective_to, period.period_code)]
    acks = {a.rule_id: a.id for a in (await db.execute(select(RuleAcknowledgement).where(RuleAcknowledgement.period_id == period.id))).scalars().all()}
    return reminders, acks

async def _gather(db, period, sales_batch_id, run_id, org_id):
    """Resolve the selected inputs into engine rows, blockers, and a fingerprint of every contributing event."""
    batch = (await db.execute(select(SalesBatch).where(SalesBatch.id == sales_batch_id, SalesBatch.period_id == period.id))).scalars().first()
    if not batch: raise HTTPException(404, 'Sales import not found for this period')
    run = await _owned_run(db, run_id, org_id)
    if run.period_id != period.id: raise HTTPException(404, 'Reconciliation run not found for this period')
    if run.status != 'succeeded': raise HTTPException(400, 'Select a succeeded reconciliation run')
    blockers = []
    if batch.status != 'committed': blockers.append('Sales import is not committed.')
    sales = await sales_record_views(db, batch.id)
    included = [r for r in sales if r['decision'] == 'reviewed' and r['validation_status'] == 'ready']
    included_ids = {r['id'] for r in included}
    pending = [r for r in sales if r['decision'] != 'excluded' and r['id'] not in included_ids]
    if pending: blockers.append(f'{len(pending)} sales row(s) are pending review.')
    # document_type is added only for notes, so invoice-only payloads stay exactly as before.
    note = lambda kind: {'document_type': kind} if kind != 'invoice' else {}
    output = [{'ref': r['raw_data']['record_id'], 'row_number': r['row_number'], 'taxable_value': r['raw_data']['taxable_value'], **note(r['raw_data']['document_type']), **{h: r['raw_data'][h] for h in HEADS}} for r in included]

    results, history = await _itc_state(db, run)
    purchases = {p.record_id: p for p in (await db.execute(select(ImportRecord).where(ImportRecord.batch_id == run.purchase_batch_id, ImportRecord.is_valid.is_(True)))).scalars().all()}
    counts = dict.fromkeys((*ITC_DECISIONS, 'undecided'), 0)
    itc = []
    for r in results:
        latest = history[r.result_id][-1].decision if history[r.result_id] else 'undecided'
        counts[latest] += 1
        if latest == 'claim':
            p = purchases[_ids(r.purchase_record_ids)[0]]
            itc.append({'ref': r.result_id, 'purchase_record_id': p.record_id, **note(p.document_type), **{h: format(getattr(p, h), '.2f') for h in HEADS}})
    if counts['undecided']: blockers.append(f"{counts['undecided']} reconciliation result(s) have no ITC decision.")
    ims = await ims_state(db, run.statement_batch_id)
    ims_conflicts = [why for r in results if history[r.result_id] and history[r.result_id][-1].decision == 'claim' and (why := ims_blocking(ims, _ids(r.statement_record_ids)))]
    if ims_conflicts: blockers.append(f'{len(ims_conflicts)} claimed invoice(s) are rejected or pending in IMS: ' + '; '.join(ims_conflicts[:5]))

    adjustments = (await db.execute(select(Adjustment).where(Adjustment.period_id == period.id).order_by(Adjustment.created_at, Adjustment.id))).scalars().all()
    voids = (await db.execute(select(AdjustmentVoid).join(Adjustment).where(Adjustment.period_id == period.id))).scalars().all()
    voided = {v.adjustment_id for v in voids}
    active = [{'ref': str(a.id), 'type': a.adjustment_type, 'head': a.tax_head, 'amount': format(a.amount, '.2f'), 'note': a.note} for a in adjustments if a.id not in voided]

    reminders, acks = await reminder_state(db, period, org_id)
    open_reminders = [r for r in reminders if r.id not in acks]
    if open_reminders: blockers.append(f'{len(open_reminders)} reminder(s) to acknowledge: ' + '; '.join(r.note[:80] for r in open_reminders))

    reviews = (await db.execute(select(SalesReview.id).join(SalesRecord).where(SalesRecord.batch_id == batch.id))).scalars().all()
    found_rules = await setoff_rules(db, org_id, period.period_code)
    fingerprint = hashlib.sha256(json.dumps({
        'engine': ENGINE_VERSION, 'sales_batch': str(batch.id), 'sales_hash': batch.file_hash, 'sales_status': batch.status,
        'sales_reviews': sorted(map(str, reviews)), 'run': str(run.id),
        'itc_decisions': sorted(str(d.id) for ds in history.values() for d in ds),
        'adjustments': sorted(str(a.id) for a in adjustments), 'voids': sorted(str(v.id) for v in voids),
        # Added only when present, so drafts saved before reminders existed keep their original fingerprint.
        **({'reminder_acks': sorted(str(a) for a in acks.values())} if acks else {}),
        # Same rule for set-off: a newly confirmed or changed rule makes earlier approvals out of date.
        **({'setoff_rules': sorted(str(r.id) for r in found_rules.values())} if found_rules else {}),
        **({'ims_actions': sorted(str(a.id) for acts in ims.values() for a in acts)} if ims else {}),
    }, sort_keys=True).encode()).hexdigest()
    try:
        result = compute(output, itc, active)
    except ValueError as exc:
        blockers.append(f'Calculation error: {exc}')
        result = None
    setoff_out = setoff_section(result, found_rules)
    rcm = {}
    for a in active:
        if a['type'] == 'rcm_liability': rcm[a['head']] = format(Decimal(rcm.get(a['head'], '0.00')) + Decimal(a['amount']), '.2f')
    outward_taxable = format(sum((sign(o) * Decimal(o['taxable_value']) for o in output), Decimal('0.00')), '.2f')
    payload = {'worksheet': result, 'setoff': setoff_out, 'gstr3b': gstr3b_summary(result, setoff_out, outward_taxable, rcm), 'inputs': {'output': output, 'itc': itc, 'adjustments': active},
               'counts': {'sales_rows': len(sales), 'sales_included': len(included), 'sales_pending': len(pending),
                          'sales_excluded': len(sales) - len(included) - len(pending), 'itc': counts},
               'sales_hash': batch.file_hash}
    return payload, blockers, fingerprint

async def _approval_state(db, period_id):
    approvals = (await db.execute(select(TaxApproval).where(TaxApproval.period_id == period_id).order_by(TaxApproval.created_at, TaxApproval.id))).scalars().all()
    reopens = {r.approval_id: r for r in (await db.execute(select(TaxReopen).join(TaxApproval).where(TaxApproval.period_id == period_id))).scalars().all()}
    return approvals, reopens

async def _draft_views(db, period, org_id):
    drafts = (await db.execute(select(TaxDraft).where(TaxDraft.period_id == period.id).order_by(TaxDraft.created_at.desc(), TaxDraft.id.desc()))).scalars().all()
    approvals, reopens = await _approval_state(db, period.id)
    by_draft = {a.draft_id: a for a in approvals}
    evidence = defaultdict(list)
    for e in (await db.execute(select(FilingEvidence).join(TaxApproval).where(TaxApproval.period_id == period.id).order_by(FilingEvidence.created_at))).scalars().all():
        evidence[e.approval_id].append({'id': str(e.id), 'arn': e.arn, 'filed_on': e.filed_on, 'note': e.note, 'actor_id': str(e.user_id), 'created_at': e.created_at.isoformat(), 'label': 'user-reported, not verified'})
    views = []
    for d in drafts:
        a = by_draft.get(d.id)
        state, approval = 'draft', None
        if a:
            reopen = reopens.get(a.id)
            if reopen: state = 'reopened'
            else:
                _, _, current = await _gather(db, period, d.sales_batch_id, d.run_id, org_id)
                state = 'approved' if current == d.fingerprint else 'approved_stale'
            approval = {'id': str(a.id), 'actor_id': str(a.user_id), 'note': a.note, 'created_at': a.created_at.isoformat(),
                        'reopen': {'reason': reopen.reason, 'actor_id': str(reopen.user_id), 'created_at': reopen.created_at.isoformat()} if reopen else None,
                        'filing_evidence': evidence[a.id]}
        views.append({'id': str(d.id), 'state': state, 'sales_batch_id': str(d.sales_batch_id), 'run_id': str(d.run_id), 'engine_version': d.engine_version,
                      'fingerprint': d.fingerprint, 'blockers': d.blockers, 'payload': d.payload, 'actor_id': str(d.user_id), 'created_at': d.created_at.isoformat(), 'approval': approval})
    return views

# --- ITC decisions --------------------------------------------------------------------------

@router.get('/runs/{run_id}/itc-decisions')
async def itc_decisions(run_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    run = await _owned_run(db, run_id, org_id)
    results, history = await _itc_state(db, run)
    # Tax on the purchase record: the credit at stake if the result were claimed (display only; negative for a credit note).
    tax = {p.record_id: sign({'document_type': p.document_type}) * (p.cgst + p.sgst + p.igst + p.cess) for p in (await db.execute(select(ImportRecord).where(ImportRecord.batch_id == run.purchase_batch_id, ImportRecord.is_valid.is_(True)))).scalars().all()}
    return [{'result_id': r.result_id, 'status': r.status, 'reason': r.reason, 'purchase_record_ids': _ids(r.purchase_record_ids),
             'itc_at_stake': format(tax[_ids(r.purchase_record_ids)[0]], '.2f') if _ids(r.purchase_record_ids) and _ids(r.purchase_record_ids)[0] in tax else None,
             'statement_record_ids': _ids(r.statement_record_ids), 'decision': history[r.result_id][-1].decision if history[r.result_id] else 'undecided',
             'previous_decision_id': str(history[r.result_id][-1].id) if history[r.result_id] else None,
             'history': [_decision_view(d) for d in history[r.result_id]]} for r in results]

class ItcItem(BaseModel):
    model_config = ConfigDict(extra='forbid')
    result_id: str
    decision: Literal['claim', 'not_claimed', 'deferred']
    previous_decision_id: UUID | None = None

class ItcInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    items: list[ItcItem] = Field(min_length=1, max_length=1000)
    note: str = Field(min_length=1, max_length=2000)

@router.post('/runs/{run_id}/itc-decisions')
async def decide_itc(run_id: UUID, data: ItcInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    if not data.note.strip(): raise HTTPException(400, 'A decision note is required')
    # Lock the run so concurrent reviewers cannot both append against the same previous decision.
    run = await _owned_run(db, run_id, org_id, lock=True)
    if run.status != 'succeeded': raise HTTPException(400, 'ITC decisions need a succeeded run')
    results, history = await _itc_state(db, run)
    ims = await ims_state(db, run.statement_batch_id)
    by_id = {r.result_id: r for r in results}
    if len({i.result_id for i in data.items}) != len(data.items): raise HTTPException(400, 'Each result may appear once per request')
    for item in data.items:
        result = by_id.get(item.result_id)
        if not result: raise HTTPException(404, f'Result {item.result_id} not found')
        if item.decision == 'claim' and result.status != 'matched':
            raise HTTPException(400, 'Only exactly matched results can be claimed in calc-v1')
        if item.decision == 'claim' and (why := ims_blocking(ims, _ids(result.statement_record_ids))):
            raise HTTPException(400, f'Cannot claim: {why}.')
        latest = history[item.result_id][-1].id if history[item.result_id] else None
        if latest != item.previous_decision_id: raise HTTPException(409, 'A newer ITC decision exists. Reload before saving.')
    for item in data.items:
        db.add(ItcDecision(run_id=run.id, result_id=item.result_id, user_id=user.id, decision=item.decision, note=data.note.strip()))
    tally = ', '.join(f'{d} {n}' for d in ITC_DECISIONS if (n := sum(i.decision == d for i in data.items)))
    audit(db, org_id, user.id, 'itc_decided', 'reconciliation_run', run.id, run.period_id, f'{len(data.items)} ITC decision(s): {tally}')
    await db.commit()
    return {'saved': len(data.items)}

# --- Adjustments ----------------------------------------------------------------------------

class AdjustmentInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    adjustment_type: Literal['rcm_liability', 'other_liability', 'itc_reversal', 'other_credit', 'opening_credit']
    tax_head: Literal['igst', 'cgst', 'sgst', 'cess']
    amount: str
    note: str = Field(min_length=1, max_length=2000)

def _adjustment_view(a, void=None):
    return {'id': str(a.id), 'adjustment_type': a.adjustment_type, 'tax_head': a.tax_head, 'amount': format(a.amount, '.2f'), 'note': a.note,
            'actor_id': str(a.user_id), 'created_at': a.created_at.isoformat(),
            'void': {'reason': void.reason, 'actor_id': str(void.user_id), 'created_at': void.created_at.isoformat()} if void else None}

@router.get('/periods/{period_id}/adjustments')
async def list_adjustments(period_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db, period_id, org_id)
    adjustments = (await db.execute(select(Adjustment).where(Adjustment.period_id == period.id).order_by(Adjustment.created_at, Adjustment.id))).scalars().all()
    voids = {v.adjustment_id: v for v in (await db.execute(select(AdjustmentVoid).join(Adjustment).where(Adjustment.period_id == period.id))).scalars().all()}
    return [_adjustment_view(a, voids.get(a.id)) for a in adjustments]

@router.post('/periods/{period_id}/adjustments')
async def add_adjustment(period_id: UUID, data: AdjustmentInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db, period_id, org_id)
    if not data.note.strip(): raise HTTPException(400, 'An adjustment note is required')
    try: amount = parse_decimal(data.amount)
    except ValueError: raise HTTPException(400, 'Amount must be a nonnegative decimal with exactly two places')
    adj = Adjustment(period_id=period.id, user_id=user.id, adjustment_type=data.adjustment_type, tax_head=data.tax_head, amount=amount, note=data.note.strip())
    db.add(adj); await db.flush()
    audit(db, org_id, user.id, 'adjustment_added', 'tax_adjustment', adj.id, period.id, f'{data.adjustment_type} {data.tax_head} {format(amount, ".2f")}')
    await db.commit(); await db.refresh(adj)
    return _adjustment_view(adj)

class VoidInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    reason: str = Field(min_length=1, max_length=2000)

@router.post('/adjustments/{adjustment_id}/void')
async def void_adjustment(adjustment_id: UUID, data: VoidInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    adj = (await db.execute(select(Adjustment).join(FilingPeriod).join(GSTRegistration).join(Client).where(Adjustment.id == adjustment_id, Client.organization_id == org_id).with_for_update(of=Adjustment))).scalars().first()
    if not adj: raise HTTPException(404, 'Adjustment not found')
    if not data.reason.strip(): raise HTTPException(400, 'A void reason is required')
    if (await db.execute(select(AdjustmentVoid).where(AdjustmentVoid.adjustment_id == adj.id))).scalars().first():
        raise HTTPException(409, 'Adjustment is already void')
    void = AdjustmentVoid(adjustment_id=adj.id, user_id=user.id, reason=data.reason.strip())
    db.add(void)
    audit(db, org_id, user.id, 'adjustment_voided', 'tax_adjustment', adj.id, adj.period_id, f'{adj.adjustment_type} {adj.tax_head} {format(adj.amount, ".2f")}')
    await db.commit()
    return _adjustment_view(adj, void)

# --- Preview, drafts, approval ----------------------------------------------------------------

@router.get('/periods/{period_id}/worksheet')
async def preview(period_id: UUID, sales_batch_id: UUID, run_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db, period_id, org_id)
    payload, blockers, fingerprint = await _gather(db, period, sales_batch_id, run_id, org_id)
    return {'notice': NOTICE, 'payload': payload, 'blockers': blockers, 'fingerprint': fingerprint, 'engine_version': ENGINE_VERSION}

class DraftInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    sales_batch_id: UUID
    run_id: UUID

@router.post('/periods/{period_id}/tax-drafts')
async def create_draft(period_id: UUID, data: DraftInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db, period_id, org_id, lock=True)
    payload, blockers, fingerprint = await _gather(db, period, data.sales_batch_id, data.run_id, org_id)
    draft = TaxDraft(period_id=period.id, sales_batch_id=data.sales_batch_id, run_id=data.run_id, user_id=user.id, engine_version=ENGINE_VERSION,
                     fingerprint=fingerprint, payload=payload, blockers=blockers)
    db.add(draft); await db.flush()
    audit(db, org_id, user.id, 'tax_draft_created', 'tax_draft', draft.id, period.id, f'{len(blockers)} blocker(s)')
    await db.commit()
    return next(v for v in await _draft_views(db, period, org_id) if v['id'] == str(draft.id))

@router.get('/periods/{period_id}/tax-drafts')
async def list_drafts(period_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db, period_id, org_id)
    return await _draft_views(db, period, org_id)

async def _owned_draft(db, draft_id, org_id):
    draft = (await db.execute(select(TaxDraft).join(FilingPeriod).join(GSTRegistration).join(Client).where(TaxDraft.id == draft_id, Client.organization_id == org_id))).scalars().first()
    if not draft: raise HTTPException(404, 'Draft not found')
    return draft

class NoteInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    note: str = Field(min_length=1, max_length=2000)

@router.post('/tax-drafts/{draft_id}/approve')
async def approve(draft_id: UUID, data: NoteInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    draft = await _owned_draft(db, draft_id, org_id)
    require_role(user, *APPROVERS)
    org = await db.get(Organization, org_id)
    if org.require_separate_approver and draft.user_id == user.id:
        raise HTTPException(403, 'Firm policy: the person who drafted this worksheet cannot approve it. Ask another reviewer.')
    # Period lock serializes approvals, reopening and draft creation for this period.
    period = await owned_period(db, draft.period_id, org_id, lock=True)
    if not data.note.strip(): raise HTTPException(400, 'An approval note is required')
    latest = (await db.execute(select(TaxDraft.id).where(TaxDraft.period_id == period.id).order_by(TaxDraft.created_at.desc(), TaxDraft.id.desc()).limit(1))).scalar_one()
    if latest != draft.id: raise HTTPException(409, 'A newer draft exists. Approve the latest draft.')
    if draft.blockers: raise HTTPException(400, 'Resolve blockers before approval: ' + ' '.join(draft.blockers))
    _, live_blockers, current = await _gather(db, period, draft.sales_batch_id, draft.run_id, org_id)
    # Blockers that appeared after the draft was saved (e.g. a newly confirmed reminder) also stop approval.
    if live_blockers: raise HTTPException(400, 'Resolve blockers before approval: ' + ' '.join(live_blockers))
    if current != draft.fingerprint: raise HTTPException(409, 'Inputs changed since this draft was computed. Create a new draft.')
    approvals, reopens = await _approval_state(db, period.id)
    if any(a.id not in reopens for a in approvals): raise HTTPException(409, 'An approval is already active for this period. Reopen it first.')
    approval = TaxApproval(draft_id=draft.id, period_id=period.id, user_id=user.id, note=data.note.strip())
    db.add(approval); await db.flush()
    audit(db, org_id, user.id, 'tax_draft_approved', 'tax_approval', approval.id, period.id, f'draft {draft.id}')
    await db.commit()
    return {'id': str(approval.id), 'draft_id': str(draft.id)}

class ReasonInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    reason: str = Field(min_length=1, max_length=2000)

async def _owned_approval(db, approval_id, org_id):
    approval = (await db.execute(select(TaxApproval).join(FilingPeriod).join(GSTRegistration).join(Client).where(TaxApproval.id == approval_id, Client.organization_id == org_id))).scalars().first()
    if not approval: raise HTTPException(404, 'Approval not found')
    return approval

@router.post('/tax-approvals/{approval_id}/reopen')
async def reopen(approval_id: UUID, data: ReasonInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    approval = await _owned_approval(db, approval_id, org_id)
    require_role(user, *APPROVERS)
    await owned_period(db, approval.period_id, org_id, lock=True)
    if not data.reason.strip(): raise HTTPException(400, 'A reopening reason is required')
    if (await db.execute(select(TaxReopen).where(TaxReopen.approval_id == approval.id))).scalars().first(): raise HTTPException(409, 'Approval is already reopened')
    db.add(TaxReopen(approval_id=approval.id, user_id=user.id, reason=data.reason.strip()))
    audit(db, org_id, user.id, 'tax_approval_reopened', 'tax_approval', approval.id, approval.period_id, data.reason.strip()[:200])
    await db.commit()
    return {'status': 'reopened'}

class EvidenceInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    arn: str = Field(min_length=1, max_length=64)
    filed_on: date
    note: str = Field(default='', max_length=2000)

@router.post('/tax-approvals/{approval_id}/filing-evidence')
async def filing_evidence(approval_id: UUID, data: EvidenceInput, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    approval = await _owned_approval(db, approval_id, org_id)
    await owned_period(db, approval.period_id, org_id, lock=True)
    if (await db.execute(select(TaxReopen).where(TaxReopen.approval_id == approval.id))).scalars().first(): raise HTTPException(409, 'This approval was reopened')
    if not data.arn.strip(): raise HTTPException(400, 'An acknowledgement reference is required')
    evidence = FilingEvidence(approval_id=approval.id, user_id=user.id, arn=data.arn.strip(), filed_on=data.filed_on.isoformat(), note=data.note.strip())
    db.add(evidence); await db.flush()
    audit(db, org_id, user.id, 'filing_evidence_recorded', 'filing_evidence', evidence.id, approval.period_id, 'User-reported filing reference recorded')
    await db.commit()
    return {'id': str(evidence.id), 'label': 'user-reported, not verified'}

@router.get('/tax-drafts/{draft_id}/export')
async def export_draft(draft_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    draft = await _owned_draft(db, draft_id, org_id)
    period = await owned_period(db, draft.period_id, org_id)
    view = next(v for v in await _draft_views(db, period, org_id) if v['id'] == str(draft.id))
    payload = draft.payload
    wb = Workbook(); wb.remove(wb.active)
    def sheet(name, headers, values):
        ws = wb.create_sheet(name); ws.append(headers)
        for row in values:
            ws.append([sanitize_value(v) for v in row])
            for cell in ws[ws.max_row]: cell.data_type = 's'
        ws.freeze_panes = 'A2'
        for cell in ws[1]: cell.font = Font(bold=True)
    approval = view['approval'] or {}
    sheet('Metadata', ['key', 'value'], [('label', NOTICE), ('draft_id', draft.id), ('period_code', period.period_code), ('state', view['state']),
          ('engine_version', draft.engine_version), ('fingerprint', draft.fingerprint), ('sales_batch_id', draft.sales_batch_id), ('sales_hash', payload.get('sales_hash')),
          ('run_id', draft.run_id), ('drafted_by', draft.user_id), ('drafted_at', draft.created_at.isoformat()),
          ('approved_by', approval.get('actor_id')), ('approved_at', approval.get('created_at')), ('approval_note', approval.get('note')),
          ('reopen_reason', (approval.get('reopen') or {}).get('reason')), ('blockers', ' '.join(draft.blockers) or 'none'),
          ('exported_at', datetime.now(timezone.utc).isoformat())])
    worksheet = payload.get('worksheet') or {'heads': {}}
    columns = ['output_tax', 'liability_adjustments', 'liability', 'itc_claimed', 'itc_reversal', 'other_credit', 'opening_credit', 'credit', 'net']
    sheet('Worksheet', ['tax_head', *columns], [(h, *[worksheet['heads'].get(h, {}).get(c, '') for c in columns]) for h in HEADS])
    so = payload.get('setoff') or {'status': 'not_computed', 'reason': 'Draft saved before set-off existed.'}
    if so['status'] == 'computed':
        rows = [('status', 'computed', '', '', '', '')]
        rows += [(f"rule: {k}", json.dumps(v['value']), f"source: {v['source_reference']}", f"from {v['effective_from']}", '', '') for k, v in so.get('rules', {}).items()]
        rows += [('utilisation', u['credit_head'].upper(), u['liability_head'].upper(), u['amount'], '', '') for u in so['utilisation']]
        rows += [('head', 'cash_before_rounding', 'rounding_difference', 'cash', 'carry_forward', '')]
        rows += [(h.upper(), so['heads'][h]['cash_before_rounding'], so['heads'][h]['rounding_difference'], so['heads'][h]['cash'], so['heads'][h]['carry_forward'], '') for h in HEADS]
        rows += [('total', so['total_cash_before_rounding'], '', so['total_cash'], so['total_carry_forward'], so.get('rounding_note', ''))]
    else:
        rows = [('status', 'not computed', so.get('reason', ''), '', '', '')]
    sheet('Set-off', ['item', 'a', 'b', 'c', 'd', 'note'], rows)
    g3 = payload.get('gstr3b') or {'status': 'not_available'}
    nc = lambda v: 'not captured' if v is None else v
    if g3.get('status') == 'available':
        g_rows = [('notice', g3['notice'], '', '', '', '', '')]
        g_rows += [('3.1 ' + r['label'], nc(r['taxable_value']), *[nc(r[h]) for h in HEADS]) for r in g3['table_3_1']]
        g_rows += [('4 ' + r['label'], '', *[nc(r[h]) for h in HEADS]) for r in g3['table_4']]
        if g3['table_6_1']:
            g_rows += [(f"6.1 {p['head'].upper()} payable {p['tax_payable']}", 'paid in cash: ' + p['paid_in_cash'], *[f"via {h.upper()} credit: {p['paid_through_itc'][h]}" for h in HEADS]) for p in g3['table_6_1']]
        else:
            g_rows += [('6.1 payment', g3.get('table_6_1_note') or 'Set-off not computed', '', '', '', '', '')]
    else:
        g_rows = [('status', 'not available', '', '', '', '', '')]
    sheet('GSTR-3B view', ['table', 'taxable_value', *HEADS], g_rows)
    sheet('Output Rows', ['record_id', 'source_row', *HEADS], [(r['ref'], r['row_number'], *[r[h] for h in HEADS]) for r in payload['inputs']['output']])
    sheet('ITC Claimed', ['result_id', 'purchase_record_id', *HEADS], [(r['ref'], r['purchase_record_id'], *[r[h] for h in HEADS]) for r in payload['inputs']['itc']])
    sheet('Adjustments', ['adjustment_id', 'type', 'tax_head', 'amount', 'note'], [(a['ref'], a['type'], a['head'], a['amount'], a['note']) for a in payload['inputs']['adjustments']])
    sheet('Filing Evidence', ['arn', 'filed_on', 'note', 'label', 'recorded_at'], [(e['arn'], e['filed_on'], e['note'], e['label'], e['created_at']) for e in approval.get('filing_evidence', [])])
    stream = BytesIO(); wb.save(stream)
    return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="Synthetic_Tax_Worksheet_{draft.id}.xlsx"'})

# --- Audit history ----------------------------------------------------------------------------

@router.get('/audit-events')
async def audit_events(period_id: UUID | None = None, offset: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=200),
                       org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    query = (select(AuditEvent, User.email, FilingPeriod.period_code, Client.name)
             .outerjoin(User, AuditEvent.user_id == User.id)
             .outerjoin(FilingPeriod, AuditEvent.period_id == FilingPeriod.id)
             .outerjoin(GSTRegistration, FilingPeriod.registration_id == GSTRegistration.id)
             .outerjoin(Client, GSTRegistration.client_id == Client.id)
             .where(AuditEvent.organization_id == org_id))
    if period_id: query = query.where(AuditEvent.period_id == period_id)
    rows = (await db.execute(query.order_by(AuditEvent.created_at.desc(), AuditEvent.id.desc()).offset(offset).limit(limit))).all()
    return [{'id': str(e.id), 'action': e.action, 'resource_type': e.resource_type, 'resource_id': e.resource_id, 'summary': e.summary,
             'actor': email, 'period_id': str(e.period_id) if e.period_id else None, 'period_code': code, 'client_name': client,
             'created_at': e.created_at.isoformat()} for e, email, code, client in rows]


# --- Credit carry-forward (Phase 5B follow-up) ---------------------------------------------------

def _previous_code(code: str) -> str:
    y, m = map(int, code.split('-'))
    return f'{y - 1}-12' if m == 1 else f'{y}-{m - 1:02d}'

async def _carry_forward(db, period, org_id):
    """Closing credit from the previous month's active, current approval, and any opening credit already entered."""
    existing = [a for a in (await db.execute(select(Adjustment).where(Adjustment.period_id == period.id, Adjustment.adjustment_type == 'opening_credit'))).scalars().all()
                if not (await db.execute(select(AdjustmentVoid.id).where(AdjustmentVoid.adjustment_id == a.id))).first()]
    base = {'existing': [_adjustment_view(a) for a in existing]}
    prev_code = _previous_code(period.period_code)
    prev = (await db.execute(select(FilingPeriod).where(FilingPeriod.registration_id == period.registration_id, FilingPeriod.period_code == prev_code))).scalars().first()
    if not prev: return base | {'status': 'not_available', 'from_period': prev_code, 'reason': f'No {prev_code} period for this registration.'}
    approved = next((d for d in await _draft_views(db, prev, org_id) if d['approval'] and not d['approval']['reopen']), None)
    if not approved: return base | {'status': 'not_available', 'from_period': prev_code, 'reason': f'{prev_code} has no active approval yet.'}
    if approved['state'] == 'approved_stale': return base | {'status': 'not_available', 'from_period': prev_code, 'reason': f'The {prev_code} approval is out of date; re-approve it first.'}
    so = approved['payload'].get('setoff')
    if not so or so.get('status') != 'computed':
        return base | {'status': 'not_available', 'from_period': prev_code, 'reason': f'The approved {prev_code} draft has no computed set-off (no CA-confirmed utilisation order at the time), so its closing credit is unknown.'}
    amounts = {h: v['carry_forward'] for h, v in so['heads'].items() if Decimal(v['carry_forward']) > 0}
    status = 'already_entered' if existing else ('available' if amounts else 'nothing_to_carry')
    return base | {'status': status, 'from_period': prev_code, 'draft_id': approved['id'], 'amounts': amounts}

@router.get('/periods/{period_id}/carry-forward')
async def carry_forward(period_id: UUID, org_id: UUID = Depends(get_active_organization), db: AsyncSession = Depends(get_db)):
    return await _carry_forward(db, await owned_period(db, period_id, org_id), org_id)

@router.post('/periods/{period_id}/carry-forward')
async def apply_carry_forward(period_id: UUID, org_id: UUID = Depends(get_active_organization), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    period = await owned_period(db, period_id, org_id, lock=True)  # serialises against a second click
    cf = await _carry_forward(db, period, org_id)
    if cf['status'] == 'already_entered': raise HTTPException(409, 'An opening credit balance is already entered for this period. Void it first to bring forward again.')
    if cf['status'] != 'available': raise HTTPException(400, cf.get('reason') or 'Nothing to bring forward.')
    note = f"Brought forward from {cf['from_period']} approved draft {cf['draft_id'][:8]} (closing credit after set-off)."
    for head, amount in cf['amounts'].items():
        db.add(Adjustment(period_id=period.id, user_id=user.id, adjustment_type='opening_credit', tax_head=head, amount=Decimal(amount), note=note))
    audit(db, org_id, user.id, 'credit_brought_forward', 'filing_period', period.id, period.id, '; '.join(f'{h.upper()} {a}' for h, a in cf['amounts'].items()))
    await db.commit()
    return await _carry_forward(db, period, org_id)
